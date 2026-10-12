"""Synthetic profile generator and renderers (EVAL.md section 2.1, task 5.1).

Every expectation is written against the brief or against the engine's own contracts
(`matcher.models`, `config/verbs.yaml`, `config/negation.yaml`), never against what the generator
happened to emit the last time it ran.

tests/eval/ deliberately has no __init__.py: with pytest's default import mode it would then be
imported as a package named `eval` and shadow the real eval/ package these tests exercise.
"""

from __future__ import annotations

import hashlib
import io
import json
import re
from collections import defaultdict
from datetime import date
from pathlib import Path

import docx
import pymupdf
import pytest
from pydantic import ValidationError

from eval.generate import bullets, corpus, doc, lexicon, postings, profiles, render, vocabulary
from matcher.errors import Flag
from matcher.graph.model import SkillGraph, SkillNode
from matcher.models import (
    DatePrecision,
    DegreeLevel,
    Evidence,
    Necessity,
    Polarity,
    Resume,
    SkillKind,
    SkillSource,
    Track,
    VerbClass,
)
from matcher.score.evidence import BulletMention, analyze_bullet, is_quantified

AS_OF = profiles.GENERATION_AS_OF

# Tiers that mean "shown in the work history" (a skills-blob or certification line is not).
WORK_TIERS = {Evidence.ROLE_BULLET, Evidence.ROLE_BULLET_QUANTIFIED, Evidence.ROLE_TITLE}
TIER_RANK = {
    Evidence.SKILLS_BLOB: 0,
    Evidence.ROLE_BULLET: 1,
    Evidence.ROLE_BULLET_QUANTIFIED: 2,
    Evidence.ROLE_TITLE: 3,
    Evidence.CERTIFICATION: 4,
}

EXPECTED_TEMPLATES = (
    "chronological",
    "functional",
    "two_column",
    "academic_cv",
    "ats_plain",
    "designer_heavy",
    "federal_long",
    "military_transition",
)
EXPECTED_FORMATS = ("pdf_latex", "pdf_word", "pdf_canva_like", "docx", "txt", "pdf_scanned_300dpi")


# --------------------------------------------------------------------------------------------
# fixtures
# --------------------------------------------------------------------------------------------


@pytest.fixture(scope="module")
def vocab() -> vocabulary.Vocabulary:
    return vocabulary.default_vocabulary()


@pytest.fixture(scope="module")
def posts(vocab):
    return postings.generate_postings(vocab, n_per_family=2, seed=11)  # 16 postings


@pytest.fixture(scope="module")
def profs(vocab, posts):
    return profiles.generate_profiles(vocab, posts, 128, seed=11)  # 8 per posting, 2 of each level


@pytest.fixture(scope="module")
def posting_by_id(posts):
    return {postings.posting_id(p): p for p in posts}


@pytest.fixture(scope="module")
def sample_profile(profs):
    """The smallest level-3 software profile: quick to render 48 times, yet a full career."""
    candidates = [p for p in profs if p.fit_level == 3 and p.family == "software"]
    assert candidates
    return min(candidates, key=lambda p: sum(len(r.bullets) for r in p.roles))


def _norm(text: str) -> str:
    return " ".join(text.split())


# --------------------------------------------------------------------------------------------
# vocabulary
# --------------------------------------------------------------------------------------------


def test_default_vocabulary_shape(vocab) -> None:
    assert 140 <= len(vocab.skills) <= 175  # "~150 skills"
    assert len(vocab.families) == 8
    ids = [s.id for s in vocab.skills]
    assert len(set(ids)) == len(ids)
    assert all(re.fullmatch(r"skill\.[a-z][a-z0-9_]*", i) for i in ids)
    assert {s.kind for s in vocab.skills} == set(SkillKind)
    for s in vocab.skills:
        assert s.families and set(s.families) <= set(vocab.families)
    # fields of study are separate from skills, and ids stay unique across both
    assert vocab.fields
    all_ids = ids + [f.id for f in vocab.fields]
    assert len(set(all_ids)) == len(all_ids)


def test_default_vocabulary_forms_are_unambiguous(vocab) -> None:
    """Two skills must never share a spelling, and no spelling may carry a digit (a digit in a
    non-quantified bullet would make the quantified flag ambiguous)."""
    seen: dict[str, str] = {}
    for s in (*vocab.skills, *vocab.fields):
        for form in (s.label, *s.aliases):
            key = form.casefold()
            assert key not in seen, f"{form!r} is spelled by both {seen[key]} and {s.id}"
            seen[key] = s.id
            assert not re.search(r"\d", form), form
            assert form.isascii(), form
            assert len(form) >= 2, form


def test_every_family_can_staff_a_posting(vocab) -> None:
    hard = {SkillKind.TECHNOLOGY, SkillKind.TOOL, SkillKind.METHODOLOGY, SkillKind.DOMAIN}
    for family in vocab.families:
        in_family = [s for s in vocab.skills if family in s.families]
        assert sum(1 for s in in_family if s.kind in hard) >= 20, family
        assert any(s.kind is SkillKind.SOFT for s in in_family), family
        assert any(f for f in vocab.fields if family in f.families), family


def test_find_mentions_longest_match_and_word_boundaries(vocab) -> None:
    text = "Built APIs in Java and JavaScript, plus mysql tuning, with node.js and C++ on Linux"
    found = vocabulary.find_mentions(vocab, text)
    assert [m.skill_id for m in found] == [
        "skill.java",
        "skill.javascript",
        "skill.mysql",
        "skill.nodejs",
        "skill.cpp",
        "skill.linux",
    ]
    assert [m.text for m in found][3] == "node.js"  # the surface form is kept as written
    assert all(text[m.start : m.end] == m.text for m in found)
    # aliases resolve, and a skill is never found inside a longer word
    assert [m.skill_id for m in vocabulary.find_mentions(vocab, "Postgres and Mongo")] == [
        "skill.postgresql",
        "skill.mongodb",
    ]
    assert vocabulary.find_mentions(vocab, "Prepared a summary of the SQLite results") == ()


def test_vocabulary_from_graph() -> None:
    nodes = {
        "skill.programming": SkillNode(
            id="skill.programming", label="Programming", kind=SkillKind.TECHNOLOGY
        ),
        "skill.python": SkillNode(
            id="skill.python",
            label="Python",
            kind=SkillKind.TECHNOLOGY,
            parents=("skill.programming",),
            aliases=("CPython",),
        ),
        "skill.baking": SkillNode(id="skill.baking", label="Baking", kind=SkillKind.DOMAIN),
        "skill.sourdough": SkillNode(
            id="skill.sourdough", label="Sourdough", kind=SkillKind.METHODOLOGY, parents=("skill.baking",)
        ),
    }
    graph = SkillGraph(nodes, {}, {}, "test")
    built = vocabulary.vocabulary_from_graph(graph, ["skill.sourdough", "skill.python", "skill.sourdough"])
    assert [s.id for s in built.skills] == [
        "skill.sourdough",
        "skill.python",
    ]  # order kept, duplicates dropped
    python = built.skills[1]
    assert (python.label, python.kind, python.aliases) == ("Python", SkillKind.TECHNOLOGY, ("CPython",))
    assert python.families == ("programming",)  # the root of its parent chain
    assert built.skills[0].families == ("baking",)
    assert built.families == ("baking", "programming")
    with pytest.raises(ValueError, match="skill.nope"):
        vocabulary.vocabulary_from_graph(graph, ["skill.python", "skill.nope"])


def test_generator_runs_end_to_end_on_a_vocabulary_built_from_a_graph() -> None:
    """The real taxonomy will replace the default vocabulary; nothing may depend on its content."""
    nodes = {
        "skill.alpha": SkillNode(id="skill.alpha", label="Alpha", kind=SkillKind.DOMAIN),
        "skill.beta": SkillNode(id="skill.beta", label="Beta", kind=SkillKind.DOMAIN),
    }
    kinds = (SkillKind.TECHNOLOGY, SkillKind.TOOL, SkillKind.METHODOLOGY, SkillKind.DOMAIN)
    names = {
        "skill.alpha": [
            f"{a} {n}"
            for a in ("Amber", "Birch", "Cedar", "Dune", "Ember", "Fjord")
            for n in ("Glazing", "Weaving", "Casting", "Carving")
        ],
        "skill.beta": [
            f"{a} {n}"
            for a in ("Garnet", "Hazel", "Indigo", "Jade", "Kelp", "Lotus")
            for n in ("Mapping", "Tuning", "Binding", "Forging")
        ],
    }
    ids = []
    for root, labels in names.items():
        for k, label in enumerate(labels):
            node_id = f"skill.{label.lower().replace(' ', '_')}"
            nodes[node_id] = SkillNode(id=node_id, label=label, kind=kinds[k % 4], parents=(root,))
            ids.append(node_id)
    built = vocabulary.vocabulary_from_graph(SkillGraph(nodes, {}, {}, "test"), ids)
    assert built.families == ("alpha", "beta") and not built.fields
    posts = postings.generate_postings(built, 2, seed=1)
    assert len(posts) == 4
    people = profiles.generate_profiles(built, posts, 16, seed=1)
    assert {p.fit_level for p in people} == {0, 1, 2, 3}
    for p in people:
        resume = profiles.expected_resume(p, built, AS_OF)
        assert resume.education and all(not e.field_ids for e in resume.education)
        shown = {sid for sid, *_rest in p.skills}
        for template in ("chronological", "functional", "military_transition"):
            text = render.render(p, template, "txt", built, seed=2).decode("utf-8")
            assert {m.skill_id for m in vocabulary.find_mentions(built, text)} == shown, (
                p.profile_id,
                template,
            )
    assert len(render.render(people[3], "designer_heavy", "pdf_canva_like", built)) > 1000


# --------------------------------------------------------------------------------------------
# postings
# --------------------------------------------------------------------------------------------


def test_postings_shape_and_determinism(vocab, posts) -> None:
    assert len(posts) == len(vocab.families) * 2
    assert postings.generate_postings(vocab, 2, seed=11) == posts
    assert postings.generate_postings(vocab, 2, seed=12) != posts
    assert {p.job_family for p in posts} == set(vocab.families)
    ids = {s.id: s for s in vocab.skills}
    for p in posts:
        assert p.title_raw and p.seniority is not None and p.min_years_total is not None
        skill_ids = [r.skill_id for r in p.requirements]
        assert len(set(skill_ids)) == len(skill_ids)
        assert set(skill_ids) <= set(ids)
        assert p.requirement_token_count == len(p.requirements)
        assert p.education_requirement is not None
        for r in p.requirements:
            if r.necessity is Necessity.REQUIRED:
                assert ids[r.skill_id].kind is not SkillKind.SOFT
                assert ids[r.skill_id].kind is not SkillKind.CREDENTIAL


def test_postings_carry_required_preferred_and_implied(posts) -> None:
    for p in posts:
        by_necessity = defaultdict(list)
        for r in p.requirements:
            by_necessity[r.necessity].append(r)
        assert 8 <= len(by_necessity[Necessity.REQUIRED]) <= 12
        assert 3 <= len(by_necessity[Necessity.PREFERRED]) <= 5
        assert 2 <= len(by_necessity[Necessity.IMPLIED]) <= 4
        for r in p.requirements:
            if r.min_years is not None:
                assert r.min_years <= max(1, p.min_years_total)


def test_posting_ids_are_unique_and_stable(posts, vocab) -> None:
    ids = [postings.posting_id(p) for p in posts]
    assert len(set(ids)) == len(ids)
    again = postings.generate_postings(vocab, 2, seed=11)
    assert [postings.posting_id(p) for p in again] == ids
    assert all(re.fullmatch(r"[a-z_]+-[0-9a-f]{8}", i) for i in ids)


# --------------------------------------------------------------------------------------------
# candidate profiles: construction of the fit levels
# --------------------------------------------------------------------------------------------


def _required(posting) -> set[str]:
    return {r.skill_id for r in posting.requirements if r.necessity is Necessity.REQUIRED}


def _coverage(profile, posting) -> float:
    required = _required(posting)
    shown = {sid for sid, _years, tier, _last in profile.skills if tier in WORK_TIERS}
    return len(required & shown) / len(required)


def test_profiles_are_deterministic_and_prefix_stable(vocab, posts, profs) -> None:
    again = profiles.generate_profiles(vocab, posts, 128, seed=11)
    assert again == profs
    assert [p.model_dump_json() for p in again] == [p.model_dump_json() for p in profs]
    assert profiles.generate_profiles(vocab, posts, 128, seed=12) != profs
    assert (
        profiles.generate_profiles(vocab, posts, 20, seed=11) == profs[:20]
    )  # profile i depends on (seed, i)


def test_fit_levels_spread_over_every_posting(posts, profs, posting_by_id) -> None:
    assert len({p.profile_id for p in profs}) == len(profs)
    levels = defaultdict(set)
    for p in profs:
        assert p.target_posting_id in posting_by_id
        assert 0 <= p.fit_level <= 3
        levels[p.target_posting_id].add(p.fit_level)
    assert set(levels) == set(posting_by_id)  # every posting is somebody's target
    assert all(found == {0, 1, 2, 3} for found in levels.values())


def test_fit_level_coverage_is_constructed(profs, posting_by_id) -> None:
    """EVAL 2.1 / brief: level 3 covers >= 90 % of the REQUIRED skills with work-history evidence,
    level 2 about 70 %, level 1 about 40 %, level 0 is a different family."""
    by_level = defaultdict(list)
    for p in profs:
        by_level[p.fit_level].append(_coverage(p, posting_by_id[p.target_posting_id]))
    assert min(by_level[3]) >= 0.9
    assert all(0.6 <= c <= 0.8 for c in by_level[2]), sorted(by_level[2])
    assert all(0.3 <= c <= 0.5 for c in by_level[1]), sorted(by_level[1])
    assert all(c == 0.0 for c in by_level[0])
    means = {level: sum(v) / len(v) for level, v in by_level.items()}
    assert means[3] > means[2] > means[1] > means[0]


def test_a_posting_with_one_required_skill_staffs_every_level(vocab, posts) -> None:
    # Regression (Task 7.1 part B, real postings): levels 1 and 2 cover at least one required skill,
    # so a one-skill posting left none to "stuff" on the skills line and random.sample raised.
    posting = posts[0]
    first = next(r for r in posting.requirements if r.necessity is Necessity.REQUIRED)
    one = posting.model_copy(update={"requirements": tuple(
        r for r in posting.requirements if r is first or r.necessity is not Necessity.REQUIRED)})
    people = profiles.generate_profiles(vocab, [one], 40, seed=3)  # 10 per level: the 30 % branch is hit
    assert {p.fit_level for p in people} == {0, 1, 2, 3}
    for p in people:
        if p.fit_level:
            assert _coverage(p, one) == 1.0  # the one required skill is shown in the work history


def test_level_zero_is_a_different_family(profs, posting_by_id) -> None:
    for p in profs:
        posting = posting_by_id[p.target_posting_id]
        if p.fit_level == 0:
            assert p.family != posting.job_family
        else:
            assert p.family == posting.job_family


def test_level_three_meets_years_seniority_and_education(vocab, profs, posting_by_id) -> None:
    checked = 0
    for p in profs:
        if p.fit_level != 3:
            continue
        posting = posting_by_id[p.target_posting_id]
        resume = profiles.expected_resume(p, vocab, AS_OF)
        assert resume.total_experience_months is not None
        assert resume.total_experience_months / 12 >= posting.min_years_total
        latest = resume.roles[0].seniority
        assert latest.track is posting.seniority.track and latest.level >= posting.seniority.level
        required_level = posting.education_requirement.min_level
        assert max(e.level.rank for e in resume.education) >= required_level.rank
        checked += 1
    assert checked >= 30


def test_lower_levels_are_visibly_weaker_on_years_and_seniority(vocab, profs, posting_by_id) -> None:
    gap = defaultdict(list)
    for p in profs:
        if p.fit_level == 0:
            continue
        posting = posting_by_id[p.target_posting_id]
        resume = profiles.expected_resume(p, vocab, AS_OF)
        gap[p.fit_level].append(resume.total_experience_months / 12 - posting.min_years_total)
    mean = {level: sum(v) / len(v) for level, v in gap.items()}
    assert mean[3] > mean[2] > mean[1]


# --------------------------------------------------------------------------------------------
# candidate profiles: the ground truth is clean
# --------------------------------------------------------------------------------------------


def _scrub(vocab, text: str) -> str:
    """The text with every skill mention blanked: a skill's own spelling ("Design Systems", "Lead
    Generation") is a name, not prose."""
    for m in reversed(vocabulary.find_mentions(vocab, text)):
        text = text[: m.start] + " " + text[m.end :]
    return text


def _negation_terms(config) -> list[str]:
    return [t.lower() for t in (*config.negation.hard_negation.terms, *config.negation.aspirational.terms)]


def _all_bullets(profile):
    for role in profile.roles:
        yield from role.bullets
    for project in profile.projects:
        yield from project.bullets


def _read(vocab, config, text: str):
    """What the engine makes of a bullet: one `MentionEvidence` per skill the vocabulary finds."""
    mentions = [
        BulletMention(canonical_id=m.skill_id, start=m.start, end=m.end)
        for m in vocabulary.find_mentions(vocab, text)
    ]
    return analyze_bullet(text, mentions, config)


ALL_PROFESSIONS = (*lexicon.FAMILY_DATA.values(), lexicon.GENERIC_FAMILY)


def test_phrase_banks_spell_no_skill_negation_or_stray_digit(vocab, config) -> None:
    """The banks bullets are built from must not smuggle in a skill, a negation or aspiration
    term, or a number: each bullet then has exactly the mentions and the polarity its ground truth
    claims. Which verb a sentence reads as is the engine's call, tested against the engine below."""
    negations = _negation_terms(config)

    def problems(text: str) -> list[str]:
        padded = " " + re.sub(r"[^a-z]+", " ", _scrub(vocab, text).lower()) + " "
        found = []
        if vocabulary.find_mentions(vocab, text):
            found.append("spells a skill")
        if any(f" {term} " in padded for term in negations):
            found.append("negation or aspiration")
        if not text.isascii():
            found.append("not ascii")
        return found

    plain: list[str] = []  # phrases without a skill, a negation or a digit
    for fam in ALL_PROFESSIONS:
        plain += [*fam.objects, *fam.metrics]
    plain += [*lexicon.TEAMS, *lexicon.PLAIN_TAILS, *lexicon.INSTRUMENT_JOINERS]
    for phrase in plain:
        assert not problems(phrase), phrase
        assert not re.search(r"\d", phrase), phrase
    # the numbered tails: nothing but the number is a digit
    fam_tails = [f.scale_tail for f in ALL_PROFESSIONS]
    tails = [t for t, _ in lexicon.METRIC_TAILS] + [t for t, _ in lexicon.FIXED_TAILS] + fam_tails
    for tail in tails:
        rendered = tail.format(metric="latency", n=7)
        assert not problems(rendered), rendered
        assert len(re.findall(r"\d+", rendered)) == 1, rendered
    # every verb phrase opens with an entry of the lexicon of its own class
    entries = {cls: {v.lower() for v in c.verbs} for cls, c in config.verbs.classes.items()}
    for verb_class, phrases in lexicon.VERB_PHRASES.items():
        assert phrases
        for prefix, layout in phrases:
            assert layout in {"obj", "first", "only"}
            words = prefix.lower().split()
            assert words[0] in entries[verb_class] or " ".join(words[:2]) in entries[verb_class], prefix
            assert not problems(prefix), prefix
    assert set(lexicon.VERB_PHRASES) == set(VerbClass) - {VerbClass.NO_VERB}
    # summaries and presentation extras: no skill is spelled, because the ground truth claims none
    extras = [
        *lexicon.SUMMARY_ADJECTIVES,
        *lexicon.SUMMARY_CLAUSES,
        lexicon.STUDENT_SUMMARY.format(adj="Detail-oriented", field="computer science"),
        *lexicon.PUB_ADJECTIVES,
        *lexicon.PUB_NOUNS,
        *lexicon.PUB_TOPICS,
        *lexicon.PUB_VENUES,
        *lexicon.TEACHING,
        *lexicon.FEDERAL_DUTIES,
        *lexicon.FEDERAL_KSA,
        *lexicon.MILITARY_BRANCHES,
        *lexicon.MILITARY_RANKS,
        *lexicon.SUPERVISOR_ROLES,
        *lexicon.KIND_THEMES.values(),
        *lexicon.PROJECT_NAMES,
    ]
    for fam in lexicon.FAMILY_DATA.values():
        for _code, title, translation in fam.mos:
            extras += [title, translation]
    for phrase in extras:
        assert not problems(phrase), phrase


# Skills of each kind of bullet, spelled as they are in the vocabulary: plain names, and names with
# punctuation, a hyphen, a slash or a space at either edge (".NET", "C#", "C++", "Node.js",
# "scikit-learn", "CI/CD", "A/B Testing", "Go-to-Market Strategy", "Month-End Close") that must not
# change how the engine splits a bullet into clauses.
FORMS_BY_GROUP = {
    "instrument": (
        ("Python", "SQL", "Airflow"),
        ("C++", "Node.js", "Go language"),
        (".NET", "C#", "scikit-learn"),
        ("Git",),
    ),
    "method": (("Agile",), ("A/B Testing",), ("CI/CD",), ("Go-to-Market Strategy",)),
    "domain": (("Machine Learning",), ("Month-End Close",)),
}


def test_every_sentence_the_templates_can_write_reads_as_declared_by_the_engine(vocab, config) -> None:
    """The ground truth follows the engine's current reading rules (SPEC 15.1, DECISIONS D-34a..f),
    not a word list: every sentence `compose_bullet` can produce, for every profession, verb class
    and kind of skill, goes through `analyze_bullet` and `is_quantified`. A change of the lexicon
    that makes a template read differently fails here, naming the sentence."""
    checked = 0
    for fam in ALL_PROFESSIONS:
        for verb_class in VerbClass:
            for group, form_sets in FORMS_BY_GROUP.items():
                for forms in form_sets:
                    wanted = [vocabulary.resolve_form(vocab, form) for form in forms]
                    for core in bullets.iter_cores(fam, verb_class, group, forms):
                        evidence = _read(vocab, config, core)
                        assert [e.canonical_id for e in evidence] == wanted, core
                        for e in evidence:
                            assert (e.verb_class, e.polarity) == (verb_class, Polarity.AFFIRMED), (core, e)
                        assert not is_quantified(core, config), core
                        checked += 1
            # a filler bullet (the work alone) has no mention: probe it with a verbless clause holding
            # one, which takes the class of the clause before it
            for core in bullets.iter_cores(fam, verb_class, "", ()):
                probe = f"{core}, Python"
                (e,) = _read(vocab, config, probe)
                assert (e.verb_class, e.polarity) == (verb_class, Polarity.AFFIRMED), probe
                checked += 1
    assert checked > 10_000  # the check is not vacuous


def test_quantified_tails_read_as_quantified_and_plain_tails_do_not(config) -> None:
    core = "Built backend services using Python"
    for fam in ALL_PROFESSIONS:
        lo, hi = fam.scale_range
        tails = [
            t.format(metric=m, n=n) for t, (a, b) in lexicon.METRIC_TAILS for m in fam.metrics for n in (a, b)
        ]
        tails += [t.format(n=n) for t, (a, b) in lexicon.FIXED_TAILS for n in (a, b)]
        tails += [fam.scale_tail.format(n=n) for n in (lo, hi)]
        for tail in tails:
            assert is_quantified(core + tail, config), (fam.key, tail)
    for tail in ("", *lexicon.PLAIN_TAILS):
        assert not is_quantified(core + tail, config), tail


def test_bullets_are_unambiguous_ground_truth(vocab, config, profs) -> None:
    """Every bullet of 128 generated profiles, read by the engine itself."""
    classes_seen: set[VerbClass] = set()
    quantified_seen = [False, False]
    negations = _negation_terms(config)
    for p in profs:
        texts = set()
        for b in _all_bullets(p):
            assert b.text not in texts, "duplicate bullet inside one profile"
            texts.add(b.text)
            # exactly the mentions the ground truth claims, in order ...
            evidence = _read(vocab, config, b.text)
            assert [e.canonical_id for e in evidence] == list(b.skill_ids), b.text
            # ... each read with the declared verb class and neither negated nor aspirational
            for e in evidence:
                assert (e.verb_class, e.polarity) == (b.verb_class, Polarity.AFFIRMED), (b.text, e)
            # a quantified outcome exactly when the ground truth says so; by the generator's own
            # convention that is also exactly when the bullet holds a digit
            assert is_quantified(b.text, config) == b.quantified, b.text
            assert b.quantified == bool(re.search(r"\d", b.text)), b.text
            padded = " " + re.sub(r"[^a-z]+", " ", _scrub(vocab, b.text).lower()) + " "
            assert not any(f" {term} " in padded for term in negations), b.text
            assert b.text[0].isupper() and b.text.isascii()
            classes_seen.add(b.verb_class)
            quantified_seen[b.quantified] = True
    assert classes_seen == set(VerbClass)  # every class of the lexicon is exercised
    assert all(quantified_seen)


def test_free_text_carries_no_stray_skill_mentions(vocab, profs) -> None:
    for p in profs:
        assert vocabulary.find_mentions(vocab, p.summary) == ()
        for role in p.roles:
            assert role.employer and role.industry
            title_ids = [m.skill_id for m in vocabulary.find_mentions(vocab, role.title)]
            assert len(title_ids) <= 1
            bullet_ids = {sid for b in role.bullets for sid in b.skill_ids}
            assert set(title_ids) <= bullet_ids  # a title skill is always shown in the role's bullets too
        for project in p.projects:
            assert vocabulary.find_mentions(vocab, project.name) == ()


def test_dates_are_month_precision_and_well_ordered(profs) -> None:
    for p in profs:
        assert p.roles
        starts = [r.start for r in p.roles]
        assert starts == sorted(starts, reverse=True)  # most recent first
        for i, r in enumerate(p.roles):
            assert r.start.day == 1 and r.start < AS_OF
            if r.end is None:
                assert i == 0  # only the newest role can still be running
            else:
                assert r.end.day == 1 and r.start < r.end <= AS_OF
            if r.is_internship:
                assert (r.seniority.track, r.seniority.level) == (Track.IC, 0)
        for project in p.projects:
            assert project.start.day == 1 and project.start <= project.end <= AS_OF


def test_skills_tuple_agrees_with_roles_blob_and_certifications(vocab, profs) -> None:
    for p in profs:
        resume = profiles.expected_resume(p, vocab, AS_OF)
        refs = defaultdict(list)
        for role in resume.roles:
            for ref in role.skills:
                refs[ref.canonical_id].append(ref)
        for project in resume.projects:
            for ref in project.skills:
                refs[ref.canonical_id].append(ref)
        for ref in (*resume.certifications, *resume.skills_blob):
            refs[ref.canonical_id].append(ref)
        recorded = {sid: (years, tier, last) for sid, years, tier, last in p.skills}
        assert len(recorded) == len(p.skills)  # one row per skill
        assert set(recorded) == set(refs)
        for sid, (years, tier, last) in recorded.items():
            best = max((r.evidence for r in refs[sid]), key=TIER_RANK.__getitem__)
            assert tier is best, sid
            used = [r.last_used for r in refs[sid] if r.last_used is not None]
            assert last == (max(used) if used else None), sid
            if tier in (Evidence.SKILLS_BLOB, Evidence.CERTIFICATION):
                assert years == 0.0 and last is None  # nothing in the work history to date it by
            assert 0.0 <= years <= 60.0


def test_years_of_use_is_the_union_of_the_roles_that_show_the_skill(vocab, profs) -> None:
    def month_index(d: date) -> int:
        return d.year * 12 + d.month - 1

    for p in profs[:40]:
        months_by_skill = defaultdict(set)
        for role in p.roles:
            end = month_index(role.end or AS_OF)
            span = set(range(month_index(role.start), end))
            for sid in {sid for b in role.bullets for sid in b.skill_ids}:
                months_by_skill[sid] |= span
        for sid, years, _tier, _last in p.skills:
            assert years == pytest.approx(len(months_by_skill.get(sid, ())) / 12, abs=0.011), sid


def test_profiles_are_synthetic_people_only(profs) -> None:
    phone = re.compile(r"\(\d{3}\) 555-01\d\d")  # the NANP range reserved for fiction
    for p in profs:
        assert p.email.endswith("@example.com")
        assert phone.fullmatch(p.phone)
        first, last = p.name.split(" ")
        assert last in {
            "Example",
            "Sample",
            "Placeholder",
            "Specimen",
            "Testcase",
            "Mockup",
            "Demo",
            "Fictional",
            "Illustrative",
            "Prototype",
            "Hypothetical",
            "Template",
        }
        assert p.locations and p.locations[0].country in {"CA", "US"}


def test_profile_models_round_trip_and_validate(profs) -> None:
    for p in profs[:10]:
        assert profiles.CandidateProfile.model_validate_json(p.model_dump_json()) == p
    p = profs[0]
    with pytest.raises(ValidationError):
        profiles.CandidateProfile.model_validate(p.model_dump() | {"fit_level": 4})
    role = p.roles[-1]
    with pytest.raises(ValidationError):
        profiles.GroundTruthRole.model_validate(
            role.model_dump() | {"end": role.start.replace(year=role.start.year - 1)}
        )


# --------------------------------------------------------------------------------------------
# expected_resume: what a perfect parser recovers
# --------------------------------------------------------------------------------------------


def test_expected_resume_validates_and_has_one_skillref_per_bullet_mention(vocab, profs) -> None:
    for p in profs:
        resume = profiles.expected_resume(p, vocab, AS_OF)
        assert isinstance(resume, Resume)
        assert Resume.model_validate(resume.model_dump()) == resume
        assert len(resume.roles) == len(p.roles)
        for gt, role in zip(p.roles, resume.roles):
            mentions = [sid for b in gt.bullets for sid in b.skill_ids]
            bullet_refs = [r for r in role.skills if r.evidence is not Evidence.ROLE_TITLE]
            assert [r.canonical_id for r in bullet_refs] == mentions
            assert len(role.skills) - len(bullet_refs) <= 1  # at most the one title mention
        for gt, project in zip(p.projects, resume.projects):
            assert [r.canonical_id for r in project.skills] == [
                sid for b in gt.bullets for sid in b.skill_ids
            ]


def test_expected_resume_maps_the_ground_truth_exactly(vocab, profs) -> None:
    id_to_label = {s.id: s.label for s in (*vocab.skills, *vocab.fields)}
    seen_title_ref = False
    for p in profs:
        resume = profiles.expected_resume(p, vocab, AS_OF)
        for gt, role in zip(p.roles, resume.roles):
            assert role.title_raw == gt.title and role.employer == gt.employer
            assert role.seniority == gt.seniority and role.industry_codes == (gt.industry,)
            assert role.is_internship == gt.is_internship
            assert role.is_current == (gt.end is None) and (role.end is None) == (gt.end is None)
            assert role.start.precision is DatePrecision.MONTH and role.start.value == gt.start
            assert role.start.earliest == role.start.latest == gt.start
            if gt.end is not None:
                assert role.end.precision is DatePrecision.MONTH and role.end.value == gt.end
            last = gt.end or AS_OF
            bullet_iter = iter([(b, sid) for b in gt.bullets for sid in b.skill_ids])
            for ref in role.skills:
                if ref.evidence is Evidence.ROLE_TITLE:
                    seen_title_ref = True
                    assert (ref.source, ref.verb_class) == (SkillSource.ROLE, VerbClass.NO_VERB)
                    assert id_to_label[ref.canonical_id].casefold() in gt.title.casefold()
                else:
                    bullet, sid = next(bullet_iter)
                    assert ref.canonical_id == sid and ref.verb_class is bullet.verb_class
                    assert ref.evidence is (
                        Evidence.ROLE_BULLET_QUANTIFIED if bullet.quantified else Evidence.ROLE_BULLET
                    )
                    assert ref.source is SkillSource.ROLE
                    assert ref.surface_form.casefold() in bullet.text.casefold()
                assert (ref.first_used, ref.last_used) == (gt.start, last)
        for gt, project in zip(p.projects, resume.projects):
            assert project.name == gt.name and project.start.value == gt.start and project.end.value == gt.end
            assert all(r.source is SkillSource.PROJECT and r.last_used == gt.end for r in project.skills)
        assert resume.education == p.education
        assert [r.canonical_id for r in resume.certifications] == list(p.certifications)
        assert all(r.evidence is Evidence.CERTIFICATION for r in resume.certifications)
        assert all(r.source is SkillSource.CERTIFICATION for r in resume.certifications)
        # the skills-section line becomes the blob, one ref per entry, in order
        assert [r.surface_form for r in resume.skills_blob] == list(p.skills_section)
        assert all(
            r.evidence is Evidence.SKILLS_BLOB and r.source is SkillSource.SKILLS_SECTION
            for r in resume.skills_blob
        )
        assert resume.locations == p.locations
    assert seen_title_ref  # the ROLE_TITLE tier is exercised somewhere in the corpus


def test_education_ground_truth_is_consistent(vocab, profs) -> None:
    fields = {f.id: f.label for f in vocab.fields}
    levels = set()
    for p in profs:
        assert p.education
        ranks = [e.level.rank for e in p.education]
        assert ranks == sorted(ranks, reverse=True)  # highest degree first
        for e in p.education:
            levels.add(e.level)
            assert set(e.field_ids) <= set(fields)
            if e.field_ids:
                assert e.field_raw == fields[e.field_ids[0]]
        assert sum(e.in_progress for e in p.education) <= 1
    assert {DegreeLevel.BACHELOR, DegreeLevel.MASTER} <= levels


def test_total_experience_months_is_the_union_of_the_dated_roles(vocab, profs) -> None:
    def month_index(d: date) -> int:
        return d.year * 12 + d.month - 1

    overlapping = 0
    for p in profs:
        covered: set[int] = set()
        total = 0
        for role in p.roles:
            span = set(range(month_index(role.start), month_index(role.end or AS_OF)))
            total += len(span)
            covered |= span
        overlapping += total > len(covered)
        assert profiles.expected_resume(p, vocab, AS_OF).total_experience_months == float(len(covered))
    assert overlapping >= 1  # concurrent roles are part of the corpus


def test_expected_resume_current_role_uses_as_of(vocab, profs) -> None:
    current = next(p for p in profs if p.roles[0].end is None)
    later = date(2027, 3, 1)
    resume = profiles.expected_resume(current, vocab, later)
    assert resume.roles[0].is_current
    assert {r.last_used for r in resume.roles[0].skills} <= {later}


def test_expected_resume_for_the_functional_template(vocab, profs) -> None:
    """A functional résumé has no chronology: undated roles, no bullets, every skill is a blob
    entry (SPEC 14.3), and the résumé is flagged."""
    p = next(p for p in profs if p.fit_level == 3 and len(p.roles) >= 2)
    resume = profiles.expected_resume(p, vocab, AS_OF, template="functional")
    assert Flag.FUNCTIONAL_RESUME in resume.flags
    assert resume.total_experience_months is None and resume.projects == ()
    assert [(r.title_raw, r.employer) for r in resume.roles] == [(r.title, r.employer) for r in p.roles]
    assert all(r.start is None and r.end is None and not r.is_current for r in resume.roles)
    assert all(ref.evidence is Evidence.ROLE_TITLE for r in resume.roles for ref in r.skills)
    assert all(ref.evidence is Evidence.SKILLS_BLOB for ref in resume.skills_blob)
    blob_ids = [r.canonical_id for r in resume.skills_blob]
    assert len(set(blob_ids)) == len(blob_ids)
    expected = {sid for sid, _y, tier, _l in p.skills if tier is not Evidence.CERTIFICATION}
    assert set(blob_ids) == expected
    assert resume.education == p.education
    # any other template gives the plain mapping
    assert profiles.expected_resume(p, vocab, AS_OF, template="chronological") == profiles.expected_resume(
        p, vocab, AS_OF
    )


def test_expected_resume_rejects_a_skills_line_it_cannot_resolve(vocab, profs) -> None:
    broken = profs[0].model_copy(update={"skills_section": (*profs[0].skills_section, "Underwater Basketry")})
    with pytest.raises(ValueError, match="Underwater Basketry"):
        profiles.expected_resume(broken, vocab, AS_OF)


# --------------------------------------------------------------------------------------------
# render: names, validation
# --------------------------------------------------------------------------------------------


def test_template_and_format_names() -> None:
    assert render.TEMPLATES == EXPECTED_TEMPLATES
    assert render.FORMATS == EXPECTED_FORMATS
    assert {f: render.FORMAT_EXTENSIONS[f] for f in render.FORMATS} == {
        "pdf_latex": "pdf",
        "pdf_word": "pdf",
        "pdf_canva_like": "pdf",
        "docx": "docx",
        "txt": "txt",
        "pdf_scanned_300dpi": "pdf",
    }
    templates = Path(render.__file__).parent / "templates"
    for name in EXPECTED_TEMPLATES:
        assert (templates / f"{name}.j2").is_file(), name


def test_unknown_template_or_format_is_rejected(vocab, sample_profile) -> None:
    with pytest.raises(ValueError, match="template"):
        render.render(sample_profile, "fancy", "txt", vocab)
    with pytest.raises(ValueError, match="format"):
        render.render(sample_profile, "chronological", "rtf", vocab)


def test_markup_parser_builds_the_document_and_rejects_malformed_input() -> None:
    text = (
        "@name Alex Example\n@headline Engineer\n@contact email || a@example.com || *\n"
        "@section skills || Skills\n@skills x, y\n@side\n@section education || Education\n"
        "@entry Degree || School || 2018 ||\n@endside\n@runhead Alex - Page\n"
    )
    parsed = doc.parse_markup(text)
    assert (parsed.name, parsed.headline, parsed.running_head) == ("Alex Example", "Engineer", "Alex - Page")
    assert [(c.kind, c.text, c.icon) for c in parsed.contacts] == [("email", "a@example.com", "*")]
    assert [(s.kind, s.side) for s in parsed.sections] == [("skills", False), ("education", True)]
    assert parsed.has_side and parsed.sections[1].items[0].args == ("Degree", "School", "2018", "")
    assert parsed.with_sidebar(frozenset({"skills"})).sections[0].side
    assert doc.clean_field(" a || b\n  c ") == "a | b c"
    for bad in (
        "@name A\n@bullet x\n",  # a bullet outside any section
        "@name A\n@bogus x\n",
        "@name A\n@section a || b\n@bullet x || y\n",  # a field separator inside a bullet
        "@name A\n@side\n",  # never closed
        "@name A\n@endside\n",
        "@name A\nplain text\n",
        "@section a || b\n",  # no name
        "@name A\n@section a || b\n@contact email || x\n",  # a contact after a section
    ):
        with pytest.raises(ValueError):
            doc.parse_markup(bad)


# --------------------------------------------------------------------------------------------
# render: determinism
# --------------------------------------------------------------------------------------------

DETERMINISM_SAMPLE = [
    ("chronological", "pdf_latex"),
    ("two_column", "docx"),
    ("designer_heavy", "pdf_canva_like"),
    ("ats_plain", "txt"),
    ("functional", "pdf_word"),
    ("federal_long", "docx"),
    ("chronological", "pdf_scanned_300dpi"),
]


@pytest.mark.parametrize(("template", "fmt"), DETERMINISM_SAMPLE)
def test_render_is_byte_identical_for_the_same_seed(vocab, sample_profile, template, fmt) -> None:
    first = render.render(sample_profile, template, fmt, vocab, seed=5)
    second = render.render(sample_profile, template, fmt, vocab, seed=5)
    assert first == second
    assert hashlib.sha256(first).hexdigest() == hashlib.sha256(second).hexdigest()


def test_render_does_not_read_the_clock_or_global_random(vocab, sample_profile) -> None:
    import random
    import time

    random.seed(1)
    a = render.render(sample_profile, "chronological", "docx", vocab, seed=2)
    time.sleep(1.1)  # zip member timestamps and PDF dates would change here if they came from the clock
    random.seed(99)
    b = render.render(sample_profile, "chronological", "docx", vocab, seed=2)
    assert a == b


def test_the_seed_varies_the_scan_but_not_the_content(vocab, sample_profile) -> None:
    a = render.render(sample_profile, "chronological", "pdf_scanned_300dpi", vocab, seed=1)
    b = render.render(sample_profile, "chronological", "pdf_scanned_300dpi", vocab, seed=2)
    assert a != b
    texts = {s: render.render(sample_profile, "chronological", "txt", vocab, seed=s).decode() for s in (1, 2)}
    for text in texts.values():
        for role in sample_profile.roles:
            for bullet in role.bullets:
                assert bullet.text in text


# --------------------------------------------------------------------------------------------
# render: every template x format renders and opens
# --------------------------------------------------------------------------------------------


def _open(data: bytes, fmt: str) -> str:
    """Open a rendering with the library a parser would use; return its readable text."""
    if fmt == "txt":
        return data.decode("utf-8")
    if fmt == "docx":
        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.extend(p.text for p in cell.paragraphs)
        return "\n".join(parts)
    pdf = pymupdf.open(stream=data, filetype="pdf")
    assert pdf.page_count >= 1
    return "\n".join(page.get_text() for page in pdf)


@pytest.fixture(scope="module")
def rendered(vocab, sample_profile) -> dict[tuple[str, str], bytes]:
    return {
        (template, fmt): render.render(sample_profile, template, fmt, vocab, seed=3)
        for template in render.TEMPLATES
        for fmt in render.FORMATS
    }


def test_every_template_and_format_renders_non_empty_and_opens(rendered, sample_profile) -> None:
    assert len(rendered) == 48
    name = sample_profile.name.casefold()
    for (template, fmt), data in rendered.items():
        assert isinstance(data, bytes) and len(data) > 500, (template, fmt)
        text = _open(data, fmt)
        if fmt == "pdf_scanned_300dpi":
            continue  # image only: checked below
        assert name in _norm(text).casefold(), (template, fmt)
        assert sample_profile.roles[0].employer in _norm(text), (template, fmt)


def test_scanned_pdf_is_image_only_at_300_dpi(rendered) -> None:
    for template in render.TEMPLATES:
        pdf = pymupdf.open(stream=rendered[(template, "pdf_scanned_300dpi")], filetype="pdf")
        assert pdf.page_count >= 1
        for page in pdf:
            assert page.get_text().strip() == "", template  # no text layer
            images = page.get_images()
            assert len(images) == 1
            width, height = images[0][2], images[0][3]
            assert (width, height) == (2550, 3300)  # US Letter at 300 dpi
            assert abs(width / page.rect.width - 300 / 72) < 0.01
        # the page count follows the word-processor rendering it was scanned from
        assert (
            pdf.page_count == pymupdf.open(stream=rendered[(template, "pdf_word")], filetype="pdf").page_count
        )


def test_txt_contains_every_skills_section_label(rendered, sample_profile) -> None:
    assert sample_profile.skills_section
    for template in render.TEMPLATES:
        text = rendered[(template, "txt")].decode("utf-8")
        for label in sample_profile.skills_section:
            assert label in text, (template, label)


def test_txt_shows_exactly_the_skills_of_the_ground_truth(vocab, profs) -> None:
    """No template leaks a skill word the ground truth does not claim, and none drops one."""
    for p in profs[:24]:
        allowed = {sid for sid, _y, _t, _l in p.skills}
        for template in render.TEMPLATES:
            text = render.render(p, template, "txt", vocab, seed=1).decode("utf-8")
            found = {m.skill_id for m in vocabulary.find_mentions(vocab, text)}
            assert found == allowed, (p.profile_id, template, found ^ allowed)


def test_text_formats_show_the_whole_work_history(rendered, sample_profile) -> None:
    for template in render.TEMPLATES:
        if template == "functional":
            continue
        for fmt in ("txt", "docx", "pdf_word", "pdf_latex", "pdf_canva_like"):
            text = _norm(_open(rendered[(template, fmt)], fmt))
            for role in sample_profile.roles:
                assert role.title in text and role.employer in text, (template, fmt)
                for bullet in role.bullets:
                    assert bullet.text in text, (template, fmt, bullet.text)


def test_functional_template_omits_the_chronology(rendered, sample_profile) -> None:
    for fmt in ("txt", "docx", "pdf_word"):
        text = _open(rendered[("functional", fmt)], fmt)
        assert not re.search(r"\b(19|20)\d\d\b", text), fmt  # no year anywhere
        for role in sample_profile.roles:
            assert role.employer in _norm(text)
            for bullet in role.bullets:
                assert bullet.text not in _norm(text)  # accomplishments are not tied to roles
    # while the dated templates do print the years
    assert re.search(r"\b20\d\d\b", _open(rendered[("chronological", "txt")], "txt"))


def test_template_specific_content(rendered) -> None:
    txt = {t: rendered[(t, "txt")].decode("utf-8") for t in render.TEMPLATES}
    assert "Publications" in txt["academic_cv"] and "Teaching" in txt["academic_cv"]
    assert "Hours per week" in txt["federal_long"] and "Supervisor" in txt["federal_long"]
    assert "MOS" in txt["military_transition"] and "Civilian" in txt["military_transition"]
    assert "✉" in txt["designer_heavy"] and "●" in txt["designer_heavy"]  # envelope icon, skill bar dots
    assert "Publications" not in txt["chronological"] and "MOS" not in txt["chronological"]
    # federal résumés are long: more words than the neutral chronology of the same person
    assert len(txt["federal_long"].split()) > 1.4 * len(txt["chronological"].split())
    assert len(pymupdf.open(stream=rendered[("federal_long", "pdf_word")], filetype="pdf")) >= 2


def test_federal_resumes_always_run_to_several_pages(vocab, profs) -> None:
    """Even the thinnest career (one job, two bullets) fills more than a page in every PDF look."""

    def words(p) -> int:
        return len(render.render(p, "chronological", "txt", vocab).split())

    thinnest = sorted(profs, key=words)[:6]
    assert min(len(p.roles) for p in thinnest) == 1
    for p in thinnest:
        for fmt in ("pdf_latex", "pdf_word", "pdf_canva_like"):
            data = render.render(p, "federal_long", fmt, vocab, seed=1)
            assert pymupdf.open(stream=data, filetype="pdf").page_count >= 2, (p.profile_id, fmt)


def test_pdf_formats_have_their_own_look(rendered) -> None:
    def fonts(data):
        """The fonts text is actually set in (a page's resource list may hold unused ones)."""
        pdf = pymupdf.open(stream=data, filetype="pdf")
        return {
            span["font"]
            for page in pdf
            for block in page.get_text("dict")["blocks"]
            for line in block.get("lines", [])
            for span in line["spans"]
            if span["text"].strip()
        }

    latex = fonts(rendered[("chronological", "pdf_latex")])
    word = fonts(rendered[("chronological", "pdf_word")])
    canva = fonts(rendered[("chronological", "pdf_canva_like")])
    assert any("Times" in f for f in latex) and not any("Helvetica" in f for f in latex)
    assert any("Helvetica" in f for f in word) and not any("Times" in f for f in word)
    assert canva and not (canva & (latex | word))  # an embedded TrueType face of its own

    def page_width_used(data):
        page = pymupdf.open(stream=data, filetype="pdf")[0]
        xs = [b[0] for b in page.get_text("blocks")] + [b[2] for b in page.get_text("blocks")]
        return max(xs) - min(xs)

    # LaTeX `article` margins are much wider than Word's one inch
    assert page_width_used(rendered[("ats_plain", "pdf_latex")]) < page_width_used(
        rendered[("ats_plain", "pdf_word")]
    )


def test_canva_like_pdf_has_a_coloured_sidebar_with_contact_details(rendered, sample_profile) -> None:
    page = pymupdf.open(stream=rendered[("chronological", "pdf_canva_like")], filetype="pdf")[0]
    sidebar = [b for b in page.get_text("blocks") if sample_profile.email in b[4]]
    assert sidebar, "the contact details live in the sidebar"
    assert sidebar[0][2] < page.rect.width * 0.45  # the left column
    fills = [d for d in page.get_drawings() if d.get("fill") and d["rect"].height > page.rect.height * 0.9]
    assert fills and fills[0]["rect"].width < page.rect.width * 0.5


def test_docx_uses_real_headings_list_styles_and_tables(rendered) -> None:
    def parse(template):
        return docx.Document(io.BytesIO(rendered[(template, "docx")]))

    chrono = parse("chronological")
    styles = [p.style.name for p in chrono.paragraphs]
    assert "Title" in styles and "Heading 1" in styles and "List Bullet" in styles
    assert len(chrono.tables) == 0
    assert len(parse("ats_plain").tables) == 0  # "single-column, no tables"
    assert len(parse("two_column").tables) == 1  # the left column is a table cell
    two_column_cells = [c.text for c in parse("two_column").tables[0].rows[0].cells]
    assert len(two_column_cells) == 2


def test_rendering_contains_no_real_contact_details(rendered) -> None:
    address = re.compile(r"[\w.+-]+@[\w.-]+")
    for (template, fmt), data in rendered.items():
        if fmt == "pdf_scanned_300dpi":
            continue
        for found in address.findall(_open(data, fmt)):
            assert found.endswith("example.com"), (template, fmt, found)


# --------------------------------------------------------------------------------------------
# corpus CLI
# --------------------------------------------------------------------------------------------


def _run_corpus(out: Path, *extra: str) -> int:
    return corpus.main(
        [
            "--profiles",
            "2",
            "--seed",
            "5",
            "--postings-per-family",
            "1",
            "--jobs",
            "1",
            "--out",
            str(out),
            *extra,
        ]
    )


def test_corpus_defaults_match_the_brief() -> None:
    args = corpus.build_parser().parse_args([])
    assert (args.profiles, args.seed, args.out) == (500, 7, "eval/corpora/synthetic")


def test_corpus_writes_files_and_a_manifest(tmp_path: Path, vocab) -> None:
    out = tmp_path / "synthetic"
    assert _run_corpus(out, "--templates", "chronological,functional", "--formats", "txt,pdf_word") == 0
    postings_rows = [
        json.loads(line) for line in (out / "postings.jsonl").read_text(encoding="utf-8").splitlines()
    ]
    assert len(postings_rows) == len(vocab.families)
    profile_rows = [
        profiles.CandidateProfile.model_validate_json(line)
        for line in (out / "profiles.jsonl").read_text(encoding="utf-8").splitlines()
    ]
    assert [p.profile_id for p in profile_rows] == ["p0001", "p0002"]
    assert {row["posting_id"] for row in postings_rows} >= {p.target_posting_id for p in profile_rows}

    manifest = json.loads((out / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["seed"] == 5
    assert manifest["counts"] == {"profiles": 2, "postings": len(vocab.families), "files": 8}
    expected_paths = {
        f"{p.profile_id}/{t}__{f}.{ext}"
        for p in profile_rows
        for t in ("chronological", "functional")
        for f, ext in (("txt", "txt"), ("pdf_word", "pdf"))
    }
    assert set(manifest["files"]) == expected_paths
    for rel, digest in manifest["files"].items():
        assert hashlib.sha256((out / rel).read_bytes()).hexdigest() == digest
    for name in ("profiles", "postings"):
        assert manifest[f"{name}_sha256"] == hashlib.sha256((out / f"{name}.jsonl").read_bytes()).hexdigest()
    # the files are the renderings of the recorded profiles
    first = profile_rows[0]
    assert (out / first.profile_id / "chronological__txt.txt").read_bytes() == render.render(
        first, "chronological", "txt", vocab, seed=corpus.render_seed(5, first.profile_id, "chronological")
    )
    # and the loaders give back exactly what the generator made, for the evaluation runners
    posts = postings.generate_postings(vocab, 1, seed=5)
    assert corpus.read_postings(out / "postings.jsonl") == {postings.posting_id(p): p for p in posts}
    assert corpus.read_profiles(out / "profiles.jsonl") == profiles.generate_profiles(vocab, posts, 2, seed=5)


def test_corpus_is_byte_identical_for_the_same_seed(tmp_path: Path) -> None:
    subset = ("--templates", "chronological,two_column", "--formats", "txt,docx,pdf_latex")
    a, b = tmp_path / "a", tmp_path / "b"
    assert _run_corpus(a, *subset) == 0 and _run_corpus(b, *subset) == 0
    files_a = sorted(p.relative_to(a) for p in a.rglob("*") if p.is_file())
    assert files_a == sorted(p.relative_to(b) for p in b.rglob("*") if p.is_file())
    assert len(files_a) == 3 + 2 * 2 * 3  # manifest + two jsonl + 2 profiles x 2 templates x 3 formats
    for rel in files_a:
        assert (a / rel).read_bytes() == (b / rel).read_bytes(), rel
    other = tmp_path / "c"
    assert (
        corpus.main(
            [
                "--profiles",
                "2",
                "--seed",
                "6",
                "--postings-per-family",
                "1",
                "--jobs",
                "1",
                "--out",
                str(other),
                *subset,
            ]
        )
        == 0
    )
    assert (other / "manifest.json").read_bytes() != (a / "manifest.json").read_bytes()


def test_corpus_parallel_run_matches_the_serial_run(tmp_path: Path) -> None:
    subset = ("--templates", "ats_plain", "--formats", "txt,pdf_word")
    serial, parallel = tmp_path / "serial", tmp_path / "parallel"
    assert _run_corpus(serial, *subset) == 0
    assert (
        corpus.main(
            [
                "--profiles",
                "2",
                "--seed",
                "5",
                "--postings-per-family",
                "1",
                "--jobs",
                "2",
                "--out",
                str(parallel),
                *subset,
            ]
        )
        == 0
    )
    assert (serial / "manifest.json").read_bytes() == (parallel / "manifest.json").read_bytes()


def test_corpus_refuses_a_used_directory_unless_forced_and_never_deletes_foreign_files(
    tmp_path: Path,
) -> None:
    out = tmp_path / "synthetic"
    subset = ("--templates", "ats_plain", "--formats", "txt")
    assert _run_corpus(out, *subset) == 0
    stale = out / "p0099"  # what an earlier, larger run would have left behind
    stale.mkdir()
    (stale / "chronological__txt.txt").write_text("stale", encoding="utf-8")

    def snapshot() -> dict[Path, bytes]:
        return {p: p.read_bytes() for p in out.rglob("*") if p.is_file()}

    before = snapshot()
    with pytest.raises(SystemExit) as refused:  # a corpus is there already: no --force, no run
        _run_corpus(out, *subset)
    assert refused.value.code == 2 and snapshot() == before  # and nothing was touched

    assert _run_corpus(out, *subset, "--force") == 0
    assert not stale.exists()  # the stale folder is gone, the new corpus stands alone
    assert sorted(p.name for p in out.iterdir()) == [
        "manifest.json",
        "p0001",
        "p0002",
        "postings.jsonl",
        "profiles.jsonl",
    ]

    (out / "notes.txt").write_text(
        "not ours", encoding="utf-8"
    )  # --force never deletes what it did not write
    with pytest.raises(SystemExit):
        _run_corpus(out, *subset, "--force")
    assert (out / "notes.txt").read_text(encoding="utf-8") == "not ours"
    assert (out / "p0001" / "ats_plain__txt.txt").is_file()
    with pytest.raises(corpus.OutputDirectoryError):
        corpus.generate_corpus(out, profiles=2, force=True)
