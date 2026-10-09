"""Deterministic generator for tests/fixtures/edge/: the ingest edge-case catalogue (SPEC §14.1).

SPEC §14: "Build the fixture directory tests/fixtures/edge/ first; it is the real specification."
Every binary file in this directory is this module's output; the expected outcome of each one is
written down in tests/test_edge_ingest.py (EXPECTED), which also fails when a file here has no
expectation. Every person is synthetic ("Alex Example", example.com addresses, 555-01xx phones).

Regenerate from match-engine/ (needs reportlab, pymupdf, python-docx and Pillow):

    ./.venv/Scripts/python.exe tests/fixtures/edge/make_edge_fixtures.py

Determinism: reportlab runs with invariant=1, PyMuPDF saves with no_new_id, DOCX zips are
rewritten with fixed timestamps, and every "random" byte comes from a seeded random.Random, so
`build_all()` is a pure function (tests/test_edge_ingest.py builds it twice and compares).
Byte-identical output across library versions is not promised (zlib, FreeType and libjpeg
differ), which is why the committed files themselves are the fixtures. Fonts come from the
libraries (reportlab's bundled Vera, MuPDF's builtin Noto fonts), never from the system.
"""

from __future__ import annotations

import datetime as dt
import io
import random
import sys
import zipfile
from pathlib import Path
from typing import Callable, Iterable, Optional, Sequence

FIXTURE_DIR = Path(__file__).resolve().parent
PAGE_W, PAGE_H = 612.0, 792.0  # US Letter in points
MAX_TOTAL_BYTES = 3_000_000  # the brief's budget for the committed outputs

# Invisible / special characters, spelled as code points so no editor can swallow them.
ZWSP = chr(0x200B)  # zero-width space
ZWNJ = chr(0x200C)  # zero-width non-joiner
ZWJ = chr(0x200D)  # zero-width joiner
WJ = chr(0x2060)  # word joiner
BOM = chr(0xFEFF)  # zero-width no-break space
SHY = chr(0x00AD)  # soft hyphen
LIG_FF, LIG_FI, LIG_FL, LIG_FFI = chr(0xFB00), chr(0xFB01), chr(0xFB02), chr(0xFB03)
LDQ, RDQ, LSQ, RSQ = chr(0x201C), chr(0x201D), chr(0x2018), chr(0x2019)
BULLET = chr(0x2022)

NAME = "Alex Example"
CONTACT = "alex.example@example.com | 555-0100 | Toronto, ON"

Line = tuple[str, float, bool]  # (text, font size in pt, bold)

RESUME_EN: list[Line] = [
    (NAME, 20, True),
    (CONTACT, 10, False),
    ("SUMMARY", 12, True),
    ("Data analyst with four years of experience building reporting pipelines.", 10, False),
    ("EXPERIENCE", 12, True),
    ("Data Analyst, Example Corp    Jan 2020 - Present", 10, True),
    (f"{BULLET} Built reporting pipelines in Python and SQL for the finance team", 10, False),
    (f"{BULLET} Automated weekly dashboards in Tableau, saving six hours per week", 10, False),
    ("Junior Analyst, Sample Industries    Jun 2018 - Dec 2019", 10, True),
    (f"{BULLET} Cleaned and validated sales data with Excel and SQL", 10, False),
    ("EDUCATION", 12, True),
    ("BSc Statistics, Example University    2014 - 2018", 10, False),
    ("SKILLS", 12, True),
    ("Python, SQL, Tableau, Excel, Statistics", 10, False),
    ("CERTIFICATIONS", 12, True),
    ("Example Certified Data Analyst (2021)", 10, False),
]

# Keywords a stuffer hides; no fixture shows them visibly.
HIDDEN_WORDS = "Kubernetes Terraform Golang Scala Hadoop"

_BUILDERS: dict[str, Callable[[], bytes]] = {}


def fixture(name: str) -> Callable[[Callable[[], bytes]], Callable[[], bytes]]:
    def register(fn: Callable[[], bytes]) -> Callable[[], bytes]:
        if name in _BUILDERS:
            raise ValueError(f"duplicate fixture {name}")
        _BUILDERS[name] = fn
        return fn

    return register


def build_all() -> dict[str, bytes]:
    """Every fixture, {filename: bytes}, in registration order."""
    return {name: build() for name, build in _BUILDERS.items()}


def fixture_names() -> tuple[str, ...]:
    return tuple(_BUILDERS)


# ─── reportlab helpers (coordinates: `top` = distance of the baseline from the page top) ────


def _canvas(encrypt: object = None, pagesize: tuple[float, float] = (PAGE_W, PAGE_H)):
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=pagesize, invariant=1, encrypt=encrypt, pageCompression=1)
    return buf, c


def _finish(buf: io.BytesIO, c) -> bytes:
    c.showPage()
    c.save()
    return buf.getvalue()


def _font(bold: bool, family: str = "Helvetica") -> str:
    return f"{family}-Bold" if bold else family


def _width(text: str, size: float, bold: bool = False) -> float:
    from reportlab.pdfbase.pdfmetrics import stringWidth

    return stringWidth(text, _font(bold), size)


def _text(c, x: float, top: float, text: str, size: float = 10.0, bold: bool = False) -> None:
    c.setFont(_font(bold), size)
    c.drawString(x, PAGE_H - top, text)


def _lines(c, lines: Iterable[Line], x: float = 72.0, top: float = 50.0, leading: float = 1.45) -> float:
    """Draw lines downwards; returns the baseline `top` of the last one."""
    for text, size, bold in lines:
        top += size * leading
        _text(c, x, top, text, size, bold)
    return top


def _column(c, x: float, top: float, lines: Sequence[Line], leading: float, max_x1: float) -> float:
    """A column of lines that must stay left of `max_x1` (the fixture's gutter promise)."""
    for text, size, bold in lines:
        if x + _width(text, size, bold) > max_x1:
            raise AssertionError(f"column line too wide for its gutter: {text!r}")
        top += leading
        _text(c, x, top, text, size, bold)
    return top


def _simple_pdf(lines: Sequence[Line], **kwargs: object) -> bytes:
    buf, c = _canvas(**kwargs)
    _lines(c, lines)
    return _finish(buf, c)


# ─── PyMuPDF helpers ────────────────────────────────────────────────────────────────────────────


def _fitz():
    try:
        import pymupdf as fitz
    except ImportError:  # pragma: no cover - older wheels only ship the legacy name
        import fitz  # type: ignore[no-redef]
    return fitz


def _mupdf_bytes(doc) -> bytes:
    doc.subset_fonts()
    return doc.tobytes(garbage=3, deflate=True, no_new_id=True)


def _noto_lines_pdf(lines: Sequence[tuple[str, float]], *, script: int = 1) -> bytes:
    """Lines set in a MuPDF builtin Noto font (script 1 = Latin): it has glyphs for zero-width
    characters, ligatures and soft hyphens, so the text layer keeps those code points."""
    fitz = _fitz()
    doc = fitz.open()
    page = doc.new_page(width=PAGE_W, height=PAGE_H)
    page.insert_font(fontname="noto", fontbuffer=fitz.Font(script=script).buffer)
    top = 50.0
    for text, size in lines:
        top += size * 1.5
        page.insert_text((72, top), text, fontname="noto", fontsize=size)
    return _mupdf_bytes(doc)


# ─── Pillow helpers (scans and photos) ─────────────────────────────────────────────────────────


def _vera_path() -> str:
    import reportlab

    return str(Path(reportlab.__file__).resolve().parent / "fonts" / "Vera.ttf")


SCAN_LINES = [
    NAME,
    "alex.example@example.com  555-0100",
    "EXPERIENCE",
    "Data Analyst, Example Corp  2019 - 2023",
    "- Built reporting pipelines in Python and SQL",
    "EDUCATION",
    "BSc Statistics, Example University  2015 - 2019",
    "SKILLS",
    "Python, SQL, Excel, Tableau",
]


def _page_image(dpi: int, lines: Sequence[str], font_pt: float = 11.0, height_in: float = 11.0):
    from PIL import Image, ImageDraw, ImageFont

    width, height = round(8.5 * dpi), round(height_in * dpi)
    img = Image.new("L", (width, height), 255)
    draw = ImageDraw.Draw(img)
    font = ImageFont.truetype(_vera_path(), max(1, round(font_pt * dpi / 72)),
                              layout_engine=ImageFont.Layout.BASIC)
    x = y = round(0.75 * dpi)
    step = round(font_pt * 1.6 * dpi / 72)
    for text in lines:
        draw.text((x, y), text, fill=0, font=font)
        y += step
    return img


def _noisy(img, seed: int, sigma: float, salt_pepper: float):
    """Approximately Gaussian noise (sum of four uniforms) plus salt-and-pepper, seeded."""
    from PIL import Image

    rng = random.Random(seed)
    scale = sigma * 1.7320508  # the sum of 4 U(0,1) has standard deviation 1/sqrt(3)
    out = bytearray(img.tobytes())
    for i, value in enumerate(out):
        r = rng.random()
        noise = (rng.random() + rng.random() + rng.random() + rng.random() - 2.0) * scale
        if r < salt_pepper:
            out[i] = 0
        elif r > 1.0 - salt_pepper:
            out[i] = 255
        else:
            out[i] = min(255, max(0, round(value + noise)))
    return Image.frombytes("L", img.size, bytes(out))


def _encode(img, fmt: str, **kwargs: object) -> bytes:
    buf = io.BytesIO()
    img.save(buf, fmt, **kwargs)
    return buf.getvalue()


def _image_pdf(image_bytes: bytes, text_layer: Optional[Sequence[tuple[float, float, str]]] = None) -> bytes:
    """A one-page PDF whose only visible content is the image (a scan). `text_layer` adds an
    invisible (render mode 3) OCR layer, as phone scanner apps do: (x, baseline_y, text) in pt."""
    fitz = _fitz()
    doc = fitz.open()
    page = doc.new_page(width=PAGE_W, height=PAGE_H)
    page.insert_image(page.rect, stream=image_bytes)
    for x, y, text in text_layer or ():
        page.insert_text((x, y), text, fontname="helv", fontsize=11, render_mode=3)
    return _mupdf_bytes(doc)


# ─── DOCX helpers ──────────────────────────────────────────────────────────────────────────────

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_DOCX_NS = (
    f'xmlns:w="{W_NS}" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:v="urn:schemas-microsoft-com:vml"'
)


def _new_docx():
    from docx import Document

    doc = Document()
    props = doc.core_properties
    fixed = dt.datetime(2020, 1, 1, 0, 0, 0)
    props.author = NAME
    props.last_modified_by = NAME
    props.created = fixed
    props.modified = fixed
    props.last_printed = fixed
    props.revision = 1
    props.title = "Résumé"
    return doc


def _deterministic_zip(data: bytes) -> bytes:
    """Rewrite a zip with fixed member timestamps (python-docx stamps the current time)."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            member = zipfile.ZipInfo(info.filename, date_time=(2020, 1, 1, 0, 0, 0))
            member.compress_type = zipfile.ZIP_DEFLATED
            member.external_attr = 0o600 << 16
            dst.writestr(member, src.read(info.filename))
    return out.getvalue()


def _docx_bytes(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return _deterministic_zip(buf.getvalue())


def _docx_resume_body(doc, *, bullets: bool = True) -> None:
    doc.add_paragraph(NAME, style="Title")
    doc.add_paragraph(CONTACT)
    doc.add_heading("EXPERIENCE", level=1)
    doc.add_paragraph("Data Analyst, Example Corp    Jan 2020 - Present")
    style = "List Bullet" if bullets else None
    doc.add_paragraph("Built reporting pipelines in Python and SQL for the finance team", style=style)
    doc.add_paragraph("Automated weekly dashboards in Tableau, saving six hours per week", style=style)
    doc.add_heading("EDUCATION", level=1)
    doc.add_paragraph("BSc Statistics, Example University    2014 - 2018")
    doc.add_heading("SKILLS", level=1)
    doc.add_paragraph("Python, SQL, Tableau, Excel, Statistics")


def _txbx_paragraphs(lines: Sequence[str]) -> str:
    return "".join(f"<w:p><w:r><w:t xml:space=\"preserve\">{t}</w:t></w:r></w:p>" for t in lines)


def _text_box_run(lines: Sequence[str], *, anchored: bool) -> str:
    """A Word 2010+ text box run. Anchored boxes carry Word's VML fallback copy of the same
    content inside mc:AlternateContent (a reader that reads both would duplicate the text)."""
    paragraphs = _txbx_paragraphs(lines)
    graphic = (
        '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr><a:xfrm><a:off x="0" y="0"/>'
        '<a:ext cx="2200000" cy="1800000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        f"</wps:spPr><wps:txbx><w:txbxContent>{paragraphs}</w:txbxContent></wps:txbx><wps:bodyPr/>"
        "</wps:wsp></a:graphicData></a:graphic>"
    )
    if not anchored:
        return (
            f"<w:r {_DOCX_NS}><w:drawing><wp:inline distT=\"0\" distB=\"0\" distL=\"0\" distR=\"0\">"
            '<wp:extent cx="2200000" cy="1800000"/><wp:docPr id="2" name="Text Box 2"/>'
            f"{graphic}</wp:inline></w:drawing></w:r>"
        )
    return (
        f"<w:r {_DOCX_NS}><mc:AlternateContent><mc:Choice Requires=\"wps\"><w:drawing>"
        '<wp:anchor distT="0" distB="0" distL="114300" distR="114300" simplePos="0" '
        'relativeHeight="251659264" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="page"><wp:posOffset>4800000</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="page"><wp:posOffset>900000</wp:posOffset></wp:positionV>'
        '<wp:extent cx="2200000" cy="1800000"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
        '<wp:wrapSquare wrapText="bothSides"/><wp:docPr id="1" name="Text Box 1"/>'
        f"<wp:cNvGraphicFramePr/>{graphic}</wp:anchor></w:drawing></mc:Choice>"
        '<mc:Fallback><w:pict><v:shape style="position:absolute;margin-left:378pt;margin-top:71pt;'
        f'width:173pt;height:142pt"><v:textbox><w:txbxContent>{paragraphs}</w:txbxContent>'
        "</v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r>"
    )


def _w(xml: str):
    from docx.oxml import parse_xml

    return parse_xml(xml)


# ════════════════════════════════════════════════════════════════════════════════════════════
# PDF with a text layer
# ════════════════════════════════════════════════════════════════════════════════════════════


@fixture("pdf_clean_one_column.pdf")
def _pdf_clean() -> bytes:
    return _simple_pdf(RESUME_EN)


def _two_column_pdf() -> bytes:
    buf, c = _canvas()
    _text(c, 50, 60, NAME, 22, True)
    # Spans both columns, above them: read first (a "band" of its own).
    _text(c, 50, 80, "alex.example@example.com | 555-0100 | Toronto, ON | example.com/alex | "
                     "open to relocation", 10)
    left: list[Line] = [
        ("EXPERIENCE", 12, True),
        ("Data Analyst, Example Corp", 10, True),
        ("Jan 2020 - Present", 10, False),
        (f"{BULLET} Built reporting pipelines in Python and SQL for the", 10, False),
        ("finance team and two regional offices", 10, False),
        (f"{BULLET} Automated weekly dashboards in Tableau, which", 10, False),
        ("saved six hours of manual work per week", 10, False),
        ("Junior Analyst, Sample Industries", 10, True),
        ("Jun 2018 - Dec 2019", 10, False),
        (f"{BULLET} Cleaned sales data with Excel and SQL and", 10, False),
        ("documented every validation rule", 10, False),
    ]
    right: list[Line] = [
        ("EDUCATION", 12, True),
        ("BSc Statistics, Example University", 9.5, False),
        ("2014 - 2018", 9.5, False),
        ("SKILLS", 12, True),
        ("Python, SQL, Tableau, Excel, Statistics,", 9.5, False),
        ("Data Visualization and Forecasting", 9.5, False),
        ("CERTIFICATIONS", 12, True),
        ("Example Certified Data Analyst", 9.5, False),
        ("LANGUAGES", 12, True),
        ("English and French", 9.5, False),
    ]
    _column(c, 50, 110, left, leading=14.5, max_x1=310)
    _column(c, 330, 104, right, leading=13.0, max_x1=PAGE_W - 30)
    return _finish(buf, c)


@fixture("pdf_right_aligned_dates.pdf")
def _pdf_right_aligned_dates() -> bytes:
    """Titles on the left, dates right-aligned on the same rows, short bullets: the date group is
    a narrow annotation of those rows, not a second column."""
    buf, c = _canvas()
    rows: list[tuple[str, float, bool, Optional[str]]] = [
        (NAME, 20, True, None),
        (CONTACT, 10, False, None),
        ("EXPERIENCE", 12, True, None),
        ("Data Analyst, Example Corp", 10, True, "Jan 2020 - Present"),
        (f"{BULLET} Built reporting pipelines", 10, False, None),
        (f"{BULLET} Automated weekly dashboards", 10, False, None),
        ("Junior Analyst, Sample Industries", 10, True, "Jun 2018 - Dec 2019"),
        (f"{BULLET} Cleaned sales data", 10, False, None),
        ("Intern, Example Lab", 10, True, "May 2017 - Aug 2017"),
        (f"{BULLET} Wrote survey reports", 10, False, None),
        ("EDUCATION", 12, True, None),
        ("BSc Statistics, Example University", 10, False, "2014 - 2018"),
        ("SKILLS", 12, True, None),
        ("Python, SQL, Tableau, Excel", 10, False, None),
    ]
    top = 50.0
    for text, size, bold, right in rows:
        top += size * 1.45
        _text(c, 72, top, text, size, bold)
        if right:
            c.setFont(_font(False), size)
            c.drawRightString(540, PAGE_H - top, right)
    return _finish(buf, c)


@fixture("pdf_two_column.pdf")
def _pdf_two_column() -> bytes:
    return _two_column_pdf()


@fixture("pdf_three_column.pdf")
def _pdf_three_column() -> bytes:
    buf, c = _canvas()
    _text(c, 40, 56, NAME, 22, True)
    col1: list[Line] = [
        ("PROFILE", 11, True),
        ("Data analyst who turns raw", 9, False),
        ("sales data into clear weekly", 9, False),
        ("reports for the finance team.", 9, False),
        ("CONTACT", 11, True),
        ("alex.example@example.com", 9, False),
        ("555-0100, Toronto", 9, False),
    ]
    col2: list[Line] = [
        ("EXPERIENCE", 11, True),
        ("Data Analyst, Example Corp", 9, True),
        ("Jan 2020 - Present", 9, False),
        (f"{BULLET} Built pipelines in Python for", 9, False),
        ("the finance team and auditors", 9, False),
        (f"{BULLET} Automated dashboards in", 9, False),
        ("Tableau for regional managers", 9, False),
    ]
    col3: list[Line] = [
        ("EDUCATION", 11, True),
        ("BSc Statistics", 9, False),
        ("Example University, 2014 - 2018", 9, False),
        ("SKILLS", 11, True),
        ("Python, SQL, Tableau, Excel,", 9, False),
        ("Statistics and Forecasting", 9, False),
    ]
    _column(c, 40, 90, col1, leading=12.0, max_x1=215)
    _column(c, 235, 92, col2, leading=13.0, max_x1=410)
    _column(c, 430, 88, col3, leading=11.5, max_x1=PAGE_W - 20)
    return _finish(buf, c)


@fixture("pdf_canva_sidebar.pdf")
def _pdf_canva_sidebar() -> bytes:
    """Dark sidebar with WHITE text (visible: it sits on a dark fill), main column on the right."""
    buf, c = _canvas()
    c.setFillColorRGB(0.13, 0.20, 0.33)
    c.rect(0, 0, 200, PAGE_H, stroke=0, fill=1)
    c.setFillColorRGB(1, 1, 1)
    sidebar: list[Line] = [
        ("CONTACT", 11, True), ("alex.example@example.com", 8.5, False), ("555-0100", 8.5, False),
        ("Toronto, ON", 8.5, False), ("SKILLS", 11, True), ("Python", 9, False), ("SQL", 9, False),
        ("Tableau", 9, False), ("Excel", 9, False), ("LANGUAGES", 11, True), ("English", 9, False),
        ("French", 9, False),
    ]
    _column(c, 24, 70, sidebar, leading=16.0, max_x1=190)
    c.setFillColorRGB(0, 0, 0)
    _text(c, 230, 64, NAME, 24, True)
    _text(c, 230, 84, "Data Analyst", 12)
    c.setStrokeColorRGB(0.13, 0.20, 0.33)
    c.line(230, PAGE_H - 92, 580, PAGE_H - 92)
    main_top = 116.0
    main: list[tuple[str, float, bool, bool]] = [  # text, size, bold, has bullet glyph
        ("PROFILE", 12, True, False),
        ("Analyst who builds reporting pipelines and dashboards", 10, False, False),
        ("EXPERIENCE", 12, True, False),
        ("Data Analyst, Example Corp    Jan 2020 - Present", 10, True, False),
        ("Built reporting pipelines in Python and SQL", 10, False, True),
        ("Automated weekly dashboards in Tableau", 10, False, True),
        ("Junior Analyst, Sample Industries    Jun 2018 - Dec 2019", 10, True, False),
        ("Cleaned and validated sales data in Excel", 10, False, True),
        ("EDUCATION", 12, True, False),
        ("BSc Statistics, Example University    2014 - 2018", 10, False, False),
    ]
    for text, size, bold, bullet in main:
        main_top += 17
        if bullet:  # Canva draws the bullet glyph and the text as separate elements
            _text(c, 230, main_top, BULLET, size)
            _text(c, 242, main_top, text, size, bold)
        else:
            _text(c, 230, main_top, text, size, bold)
    return _finish(buf, c)


def _paged_pdf(pages: Sequence[Sequence[Line]], header: Optional[str], footer: Callable[[int, int], str]) -> bytes:
    buf, c = _canvas()
    total = len(pages)
    for number, lines in enumerate(pages, start=1):
        if header:
            _text(c, 72, 40, header, 9)
        text = footer(number, total)
        if text:
            _text(c, PAGE_W / 2 - _width(text, 9) / 2, PAGE_H - 30, text, 9)
        _lines(c, lines, top=70)
        c.showPage()
    c.save()
    return buf.getvalue()


_PAGE_BODIES: list[list[Line]] = [
    [(NAME, 20, True), (CONTACT, 10, False), ("EXPERIENCE", 12, True),
     ("Data Analyst, Example Corp    Jan 2020 - Present", 10, True),
     (f"{BULLET} Built reporting pipelines in Python and SQL", 10, False)],
    [("Junior Analyst, Sample Industries    Jun 2018 - Dec 2019", 10, True),
     (f"{BULLET} Cleaned and validated sales data with Excel", 10, False), ("EDUCATION", 12, True),
     ("BSc Statistics, Example University    2014 - 2018", 10, False)],
    [("SKILLS", 12, True), ("Python, SQL, Tableau, Excel, Statistics", 10, False),
     ("CERTIFICATIONS", 12, True), ("Example Certified Data Analyst (2021)", 10, False)],
]


@fixture("pdf_header_footer.pdf")
def _pdf_header_footer() -> bytes:
    return _paged_pdf(_PAGE_BODIES, f"{NAME} | Curriculum Vitae", lambda n, t: f"Page {n} of {t}")


@fixture("pdf_page_numbers_plain.pdf")
def _pdf_page_numbers_plain() -> bytes:
    return _paged_pdf(_PAGE_BODIES, None, lambda n, t: str(n))


@fixture("pdf_page_numbers_dashed.pdf")
def _pdf_page_numbers_dashed() -> bytes:
    return _paged_pdf(_PAGE_BODIES, None, lambda n, t: f"- {n} -")


@fixture("pdf_ligatures.pdf")
def _pdf_ligatures() -> bytes:
    return _noto_lines_pdf([
        (NAME, 18), (CONTACT, 10), ("EXPERIENCE", 12),
        (f"O{LIG_FFI}ce Manager, Example Corp    2019 - 2023", 10),
        (f"{BULLET} Reduced work{LIG_FL}ow e{LIG_FF}ort for {LIG_FI}nance sta{LIG_FF}", 10),
        ("SKILLS", 12), (f"Scheduling, {LIG_FI}ling, budgeting", 10),
    ])


@fixture("pdf_zero_width.pdf")
def _pdf_zero_width() -> bytes:
    return _noto_lines_pdf([
        (NAME, 18), (CONTACT, 10), ("EXPERIENCE", 12),
        ("Data Analyst, Example Corp    2019 - 2023", 10),
        (f"{BULLET} Built pipelines in Py{ZWSP}thon and {BOM}SQL on Kuber{ZWNJ}netes", 10),
        ("SKILLS", 12), (f"Ta{ZWJ}bleau, Ex{WJ}cel, Py{ZWSP}thon", 10),
    ])


@fixture("pdf_soft_hyphens.pdf")
def _pdf_soft_hyphens() -> bytes:
    return _noto_lines_pdf([
        (NAME, 18), (CONTACT, 10), ("SUMMARY", 12),
        (f"Experi{SHY}enced analyst responsible for the devel{SHY}", 10),
        ("opment of reporting tools for the finance team.", 10),
        ("EXPERIENCE", 12), ("Data Analyst, Example Corp    2019 - 2023", 10),
    ])


@fixture("pdf_hyphenated_lines.pdf")
def _pdf_hyphenated_lines() -> bytes:
    return _simple_pdf([
        (NAME, 18, True), (CONTACT, 10, False), ("EXPERIENCE", 12, True),
        ("Data Analyst, Example Corp    2019 - 2023", 10, True),
        (f"{BULLET} Led the devel-", 10, False),
        ("opment of internal reporting tools for the finance team", 10, False),
        (f"{BULLET} Self-taught in statistics and data visualization", 10, False),
    ])


@fixture("pdf_smart_quotes.pdf")
def _pdf_smart_quotes() -> bytes:
    return _simple_pdf([
        (NAME, 18, True), (CONTACT, 10, False), ("EXPERIENCE", 12, True),
        ("Data Analyst, Example Corp    2019 - 2023", 10, True),
        (f"{BULLET} Awarded {LDQ}Analyst of the Year{RDQ} for the {LSQ}Clean Data{RSQ} project", 10, False),
        (f"{BULLET} Rebuilt the team{RSQ}s reporting pipeline in Python", 10, False),
    ])


@fixture("pdf_bullet_glyphs.pdf")
def _pdf_bullet_glyphs() -> bytes:
    fitz = _fitz()
    glyphs = [0x2022, 0x25AA, 0x25E6, 0x25CF, 0x2023, 0x2219, 0x27A2, 0x2713, 0x2013]
    items = ["Python", "SQL", "Tableau", "Excel", "Statistics", "Reporting", "Dashboards", "Forecasting",
             "Budgeting"]
    rows = "".join(f"<div>{chr(g)} {item}</div>" for g, item in zip(glyphs, items))
    doc = fitz.open()
    page = doc.new_page(width=PAGE_W, height=PAGE_H)
    html = (f"<div><b>{NAME}</b></div><div>{CONTACT}</div><div><b>EXPERIENCE</b></div>"
            "<div>Data Analyst, Example Corp 2019 - 2023</div><div><b>SKILLS</b></div>" + rows)
    page.insert_htmlbox(fitz.Rect(72, 60, 540, 600), html, css="div {font-size: 11px; margin: 0 0 4px 0;}")
    return _mupdf_bytes(doc)


def _stuffed_pdf(draw_hidden: Callable[[object], None]) -> bytes:
    buf, c = _canvas()
    _lines(c, RESUME_EN)
    c.saveState()
    draw_hidden(c)
    c.restoreState()
    _text(c, 72, 740, "References available on request", 10)  # drawn after: must stay visible
    return _finish(buf, c)


@fixture("pdf_white_text.pdf")
def _pdf_white_text() -> bytes:
    def hide(c) -> None:
        c.setFillColorRGB(1, 1, 1)
        _text(c, 72, 690, HIDDEN_WORDS, 10)
        c.setFillGray(0.97)
        _text(c, 72, 705, "Rust Haskell", 10)
        c.setFillColorCMYK(0, 0, 0, 0.02)
        _text(c, 300, 705, "Elixir Erlang", 10)
        c.setFillGray(0.6)  # light grey is a design choice, not stuffing: stays
        _text(c, 72, 720, "Updated October 2026", 9)

    return _stuffed_pdf(hide)


@fixture("pdf_tiny_text.pdf")
def _pdf_tiny_text() -> bytes:
    def hide(c) -> None:
        _text(c, 72, 690, HIDDEN_WORDS, 1)
        _text(c, 72, 700, "Rust Haskell", 1.5)
        _text(c, 72, 720, "Portfolio at example.com/alex", 6)  # small but readable: stays

    return _stuffed_pdf(hide)


@fixture("pdf_invisible_render_mode.pdf")
def _pdf_invisible_render_mode() -> bytes:
    def hide(c) -> None:
        text = c.beginText(72, PAGE_H - 690)
        text.setFont("Helvetica", 10)
        text.setTextRenderMode(3)
        text.textLine(HIDDEN_WORDS)
        c.drawText(text)

    return _stuffed_pdf(hide)


@fixture("pdf_clip_render_modes.pdf")
def _pdf_clip_render_modes() -> bytes:
    """Render mode 7 (add to clip path only) is invisible and never reaches get_texttrace();
    mode 4 (fill, then clip) is visible, and MuPDF reports it twice (painted + clip copy)."""
    def hide(c) -> None:
        text = c.beginText(72, PAGE_H - 690)
        text.setFont("Helvetica", 10)
        text.setTextRenderMode(7)
        text.textLine(HIDDEN_WORDS)
        c.drawText(text)
        visible = c.beginText(72, PAGE_H - 710)
        visible.setFont("Helvetica", 10)
        visible.setTextRenderMode(4)
        visible.textLine("Fluent in English and French")
        c.drawText(visible)

    return _stuffed_pdf(hide)


@fixture("pdf_fake_bold.pdf")
def _pdf_fake_bold() -> bytes:
    """Old PDF printers fake bold by drawing the same text again a fraction of a point away."""
    buf, c = _canvas()
    top = _lines(c, RESUME_EN[:5])
    for dx in (0.0, 0.35, 0.7):
        _text(c, 72 + dx, top + 16, "Data Analyst, Example Corp    Jan 2020 - Present", 10)
    _lines(c, RESUME_EN[6:], top=top + 16)
    return _finish(buf, c)


@fixture("pdf_zero_opacity.pdf")
def _pdf_zero_opacity() -> bytes:
    def hide(c) -> None:
        c.setFillAlpha(0)
        _text(c, 72, 690, HIDDEN_WORDS, 10)

    return _stuffed_pdf(hide)


@fixture("pdf_offpage_text.pdf")
def _pdf_offpage_text() -> bytes:
    def hide(c) -> None:
        _text(c, -500, 400, HIDDEN_WORDS, 10)  # left of the page
        _text(c, 72, -60, "Rust Haskell", 10)  # above the page

    return _stuffed_pdf(hide)


@fixture("pdf_encrypted.pdf")
def _pdf_encrypted() -> bytes:
    from reportlab.lib import pdfencrypt

    enc = pdfencrypt.StandardEncryption("open-sesame", ownerPassword="owner-secret", strength=128)
    return _simple_pdf(RESUME_EN, encrypt=enc)


@fixture("pdf_owner_password_only.pdf")
def _pdf_owner_password_only() -> bytes:
    from reportlab.lib import pdfencrypt

    enc = pdfencrypt.StandardEncryption("", ownerPassword="owner-secret", canModify=0, canCopy=0,
                                        strength=128)
    return _simple_pdf(RESUME_EN, encrypt=enc)


@fixture("pdf_truncated.pdf")
def _pdf_truncated() -> bytes:
    full = _simple_pdf(RESUME_EN)
    return full[: int(len(full) * 0.6)]


@fixture("pdf_damaged_stream.pdf")
def _pdf_damaged_stream() -> bytes:
    """A 200-byte hole of zeros in the middle of the file: the page's compressed content stream no
    longer inflates (MuPDF: "zlib error"), though the file still opens."""
    full = _simple_pdf(RESUME_EN)
    middle = len(full) // 2
    return full[:middle] + bytes(200) + full[middle + 200:]


@fixture("pdf_truncated_eof_appended.pdf")
def _pdf_truncated_eof_appended() -> bytes:
    """Cut in half, then "%%EOF" appended so the trailer check passes: MuPDF repairs the xref and
    the content stream ends early ("premature end of data in flate filter")."""
    full = _simple_pdf(RESUME_EN)
    return full[: len(full) // 2] + b"\n%%EOF\n"


# Hand-built one-page PDFs with a plain FlateDecode content stream, to tell a stream whose data was
# fully recovered (a bad or missing Adler-32 checksum: MuPDF warns, every line is there) from one
# that lost data (fix round 2: damage warnings alone must not make a file CORRUPT).
HANDMADE_LINES = [
    NAME, "alex.example@example.com | 555-0100 | Toronto, ON", "SUMMARY",
    "Data analyst who builds reporting pipelines for finance teams.", "EXPERIENCE",
    "Data Analyst, Example Corp   2021 - 2024",
    "- Built weekly reports in SQL and Python for the sales team",
    "- Reduced reporting time by 40 percent with automated pipelines",
    "- Presented monthly results to the regional managers",
    "Junior Analyst, Sample Industries   2019 - 2021",
    "- Cleaned and validated sales data with Excel and SQL",
    "- Maintained the shared dashboard used by twelve teams",
    "Intern, Example Lab   2018 - 2019",
    "- Wrote survey reports and checked data quality",
    "EDUCATION", "BSc Statistics, Example University   2014 - 2018",
    "- Thesis on forecasting retail demand", "SKILLS",
    "Python, SQL, Excel, Tableau, statistics", "Forecasting, data cleaning, reporting",
    "CERTIFICATIONS", "Example Certified Data Analyst (2022)", "Example Cloud Practitioner (2023)",
    "LANGUAGES", "English and French", "PROJECTS",
    "- Open data dashboard for a city transit network",
    "- Survey sampling simulator in Python",
    "VOLUNTEERING", "Treasurer, Example Chess Club   2016 - 2018",
]


def _handmade_content() -> bytes:
    def escape(text: str) -> bytes:
        return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)").encode("latin-1")

    return (b"BT /F1 10 Tf 13 TL 72 760 Td\n" + b"".join(b"(" + escape(line) + b") Tj T*\n"
                                                       for line in HANDMADE_LINES) + b"ET\n")


def handmade_pdf(stream_data: bytes) -> bytes:
    """A minimal PDF (catalog, pages, one page, a FlateDecode content stream, Helvetica)."""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d /Filter /FlateDecode >>\nstream\n" % len(stream_data) + stream_data + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return bytes(out)


def _handmade_deflate() -> bytes:
    import zlib

    return zlib.compress(_handmade_content(), 9)


@fixture("pdf_recovered_checksum.pdf")
def _pdf_recovered_checksum() -> bytes:
    """Adler-32 checksum wrong; every byte decodes ("ignoring zlib error: incorrect data check")."""
    good = _handmade_deflate()
    return handmade_pdf(good[:-4] + b"\x00\x00\x00\x00")


@fixture("pdf_missing_checksum.pdf")
def _pdf_missing_checksum() -> bytes:
    """No Adler-32 checksum at all (a common writer bug): "premature end of data", no data lost."""
    return handmade_pdf(_handmade_deflate()[:-4])


@fixture("pdf_damaged_partial.pdf")
def _pdf_damaged_partial() -> bytes:
    """The last third of the deflate data zeroed: the first lines survive, the rest is garbage."""
    good = _handmade_deflate()
    keep = 2 * len(good) // 3
    return handmade_pdf(good[:keep] + bytes(len(good) - keep))


@fixture("pdf_deflate_cut_short.pdf")
def _pdf_deflate_cut_short() -> bytes:
    """The deflate data cut at 60 % (and /Length says so): the stream ends before its data does."""
    good = _handmade_deflate()
    return handmade_pdf(good[: len(good) * 6 // 10])


@fixture("pdf_garbled.pdf")
def _pdf_garbled() -> bytes:
    rng = random.Random(20261008)
    return b"%PDF-1.7\n%\xe2\xe3\xcf\xd3\n" + bytes(rng.getrandbits(8) for _ in range(3000))


@fixture("pdf_junk_prefix.pdf")
def _pdf_junk_prefix() -> bytes:
    # Some mail gateways prepend bytes; readers look for %PDF- in the first 1024 bytes.
    return b"X-Scanned-By: example gateway\r\n\r\n" + _simple_pdf(RESUME_EN)


# ════════════════════════════════════════════════════════════════════════════════════════════
# Scans, images, photos (OCR)
# ════════════════════════════════════════════════════════════════════════════════════════════


@fixture("scan_300dpi.pdf")
def _scan_300dpi() -> bytes:
    return _image_pdf(_encode(_page_image(300, SCAN_LINES), "PNG", optimize=True))


@fixture("scan_noisy_72dpi.pdf")
def _scan_noisy_72dpi() -> bytes:
    from PIL import ImageFilter

    img = _noisy(_page_image(72, SCAN_LINES, font_pt=9), seed=7, sigma=40, salt_pepper=0.05)
    img = img.filter(ImageFilter.GaussianBlur(0.6))
    return _image_pdf(_encode(img, "JPEG", quality=35))


@fixture("pdf_searchable_scan.pdf")
def _pdf_searchable_scan() -> bytes:
    """A phone-scanner PDF: the page is an image, plus an invisible OCR text layer on top. The
    invisible layer is the scanner's, not stuffing."""
    dpi, font_pt = 300, 11.0
    img = _page_image(dpi, SCAN_LINES, font_pt=font_pt)
    x = 0.75 * 72
    step = round(font_pt * 1.6 * dpi / 72) * 72 / dpi
    layer = [(x, 0.75 * 72 + font_pt + i * step, text) for i, text in enumerate(SCAN_LINES)]
    return _image_pdf(_encode(img, "PNG", optimize=True), text_layer=layer)


@fixture("pdf_blank_page.pdf")
def _pdf_blank_page() -> bytes:
    fitz = _fitz()
    doc = fitz.open()
    doc.new_page(width=PAGE_W, height=PAGE_H)
    return _mupdf_bytes(doc)


@fixture("image_resume.png")
def _image_resume_png() -> bytes:
    return _encode(_page_image(150, SCAN_LINES, height_in=5.0), "PNG", optimize=True)


@fixture("photo_no_text.jpg")
def _photo_no_text() -> bytes:
    from PIL import Image, ImageDraw, ImageFilter

    width, height = 640, 480
    img = Image.new("RGB", (width, height))
    pixels = img.load()
    for y in range(height):
        t = y / (height - 1)
        for x in range(width):
            s = x / (width - 1)
            pixels[x, y] = (round(90 + 120 * t + 20 * s), round(140 + 80 * t), round(220 - 60 * t))
    draw = ImageDraw.Draw(img)
    draw.ellipse((470, 60, 560, 150), fill=(250, 220, 120))
    draw.polygon([(0, 380), (150, 220), (300, 360), (420, 250), (640, 400), (640, 480), (0, 480)],
                 fill=(70, 110, 80))
    draw.polygon([(0, 430), (200, 330), (380, 420), (640, 360), (640, 480), (0, 480)], fill=(50, 90, 60))
    return _encode(img.filter(ImageFilter.GaussianBlur(2)), "JPEG", quality=80)


# ════════════════════════════════════════════════════════════════════════════════════════════
# Long documents
# ════════════════════════════════════════════════════════════════════════════════════════════


def _long_cv(pages: int) -> bytes:
    buf, c = _canvas()
    for number in range(1, pages + 1):
        lines: list[Line] = []
        if number == 1:
            lines += [(NAME, 20, True), (CONTACT, 10, False), ("EXPERIENCE", 12, True)]
        year = 1960 + number % 60
        lines += [
            (f"Research Analyst, Example Lab project {number:03d}    {year} - {year + 1}", 10, True),
            (f"{BULLET} Analysed survey data in R and Python for project {number:03d}", 10, False),
            (f"{BULLET} Wrote quarterly reports for the program office", 10, False),
            (f"{BULLET} Presented findings to {number % 7 + 2} partner agencies", 10, False),
        ]
        if number == pages:
            lines += [("EDUCATION", 12, True), ("PhD Statistics, Example University    1955 - 1960", 10, False)]
        _lines(c, lines, top=60)
        footer = f"Page {number} of {pages}"
        _text(c, PAGE_W / 2 - _width(footer, 9) / 2, PAGE_H - 30, footer, 9)
        c.showPage()
    c.save()
    return buf.getvalue()


@fixture("pdf_90_pages.pdf")
def _pdf_90_pages() -> bytes:
    return _long_cv(90)


@fixture("pdf_120_pages.pdf")
def _pdf_120_pages() -> bytes:
    return _long_cv(120)


# ════════════════════════════════════════════════════════════════════════════════════════════
# Wrong documents, minimal documents, reading order
# ════════════════════════════════════════════════════════════════════════════════════════════

TWO_LINE_RESUME = [
    f"{NAME} | alex.example@example.com",
    "Data Analyst, Example Corp (2020 - 2024): Python, SQL, Tableau",
]

COVER_LETTER = [
    NAME,
    "alex.example@example.com",
    "October 1, 2026",
    "Hiring Manager",
    "Example Corp",
    "Dear Hiring Manager,",
    "I am writing to apply for the Data Analyst position at Example Corp. Over the past",
    "four years I have built reporting pipelines in Python and SQL, and I would bring the",
    "same care to your finance team.",
    "I would welcome the chance to discuss how my experience can help your team.",
    "Sincerely,",
    NAME,
]

TRANSCRIPT = [
    "Example University",
    "Official Transcript",
    f"Student: {NAME}    Student ID: 000000",
    "Term: Fall 2016",
    "Course Code    Course Title                    Credits    Grade",
    "STAT 101       Introduction to Statistics      3.0        A",
    "MATH 120       Calculus I                      4.0        A-",
    "CS 110         Programming in Python           3.0        A",
    "Term GPA: 3.85    Cumulative GPA: 3.85",
    "Term: Winter 2017",
    "Course Code    Course Title                    Credits    Grade",
    "STAT 201       Probability                     3.0        A",
    "MATH 220       Linear Algebra                  3.0        B+",
    "Term GPA: 3.65    Cumulative GPA: 3.75",
]


@fixture("pdf_two_line_resume.pdf")
def _pdf_two_line_resume() -> bytes:
    return _simple_pdf([(t, 10, False) for t in TWO_LINE_RESUME])


@fixture("pdf_cover_letter.pdf")
def _pdf_cover_letter() -> bytes:
    return _simple_pdf([(t, 10.5, False) for t in COVER_LETTER])


@fixture("pdf_transcript.pdf")
def _pdf_transcript() -> bytes:
    buf, c = _canvas()
    top = 60.0
    for text in TRANSCRIPT:
        top += 15
        c.setFont("Courier", 9)
        c.drawString(60, PAGE_H - top, text)
    return _finish(buf, c)


@fixture("pdf_column_order_uncertain.pdf")
def _pdf_column_order_uncertain() -> bytes:
    """Sentences run ACROSS a gutter (a two-cell table read row by row). Column-by-column
    reading breaks them, which the sentence-continuity check must notice (tier B)."""
    buf, c = _canvas()
    _text(c, 50, 60, NAME, 20, True)
    _text(c, 50, 80, "alex.example@example.com | 555-0100", 10)
    _text(c, 50, 110, "EXPERIENCE", 12, True)
    _text(c, 50, 128, "Data Analyst, Example Corp    2020 - 2024", 10, True)
    rows = [
        (f"{BULLET} Designed the ingestion layer for", "the analytics team in Toronto and Montreal"),
        (f"{BULLET} Migrated the legacy reports to", "a cloud warehouse with automated tests"),
        (f"{BULLET} Trained six new analysts on", "the internal style guide and review process"),
        (f"{BULLET} Presented the quarterly results to", "the executive committee every quarter"),
    ]
    top = 128.0
    for left, right in rows:
        top += 15
        if 50 + _width(left, 10) > 315:
            raise AssertionError("left cell crosses the gutter")
        _text(c, 50, top, left, 10)
        _text(c, 330, top, right, 10)
    _text(c, 50, top + 30, "EDUCATION", 12, True)
    _text(c, 50, top + 46, "BSc Statistics, Example University    2014 - 2018", 10)
    return _finish(buf, c)


# ════════════════════════════════════════════════════════════════════════════════════════════
# Languages (D-70: supported = en, fr)
# ════════════════════════════════════════════════════════════════════════════════════════════

FRENCH = [
    (NAME, 20, True),
    ("alex.example@example.com | 555-0100 | Montréal, QC", 10, False),
    ("PROFIL", 12, True),
    ("Analyste de données avec quatre ans d'expérience dans la création de rapports.", 10, False),
    ("EXPÉRIENCE PROFESSIONNELLE", 12, True),
    ("Analyste de données, Exemple Inc.    janv. 2020 - aujourd'hui", 10, True),
    (f"{BULLET} Conception de tableaux de bord pour l'équipe des finances", 10, False),
    (f"{BULLET} Automatisation des rapports hebdomadaires avec Python et SQL", 10, False),
    ("Analyste junior, Société Exemple    juin 2018 - déc. 2019", 10, True),
    (f"{BULLET} Nettoyage et validation des données de ventes", 10, False),
    ("FORMATION", 12, True),
    ("Baccalauréat en statistique, Université Exemple    2014 - 2018", 10, False),
    ("COMPÉTENCES", 12, True),
    ("Python, SQL, Tableau, Excel, statistiques", 10, False),
    ("LANGUES", 12, True),
    ("Français et anglais", 10, False),
]

SPANISH = [
    (NAME, 20, True),
    ("alex.example@example.com | 555-0100 | Madrid", 10, False),
    ("PERFIL", 12, True),
    ("Analista de datos con cuatro años de experiencia en la creación de informes.", 10, False),
    ("EXPERIENCIA PROFESIONAL", 12, True),
    ("Analista de datos, Ejemplo S.A.    ene. 2020 - actualidad", 10, True),
    (f"{BULLET} Diseño de paneles para el equipo de finanzas", 10, False),
    (f"{BULLET} Automatización de los informes semanales con Python y SQL", 10, False),
    ("Analista júnior, Empresa Ejemplo    jun. 2018 - dic. 2019", 10, True),
    (f"{BULLET} Limpieza y validación de los datos de ventas", 10, False),
    ("EDUCACIÓN", 12, True),
    ("Licenciatura en Estadística, Universidad Ejemplo    2014 - 2018", 10, False),
    ("HABILIDADES", 12, True),
    ("Python, SQL, Tableau, Excel", 10, False),
    ("IDIOMAS", 12, True),
    ("Español e inglés", 10, False),
]

GERMAN = [
    (NAME, 20, True),
    ("alex.example@example.com | 555-0100 | Berlin", 10, False),
    ("PROFIL", 12, True),
    ("Datenanalyst mit vier Jahren Erfahrung in der Erstellung von Berichten.", 10, False),
    ("BERUFSERFAHRUNG", 12, True),
    ("Datenanalyst, Beispiel GmbH    01/2020 - heute", 10, True),
    (f"{BULLET} Entwicklung von Dashboards für die Finanzabteilung", 10, False),
    (f"{BULLET} Automatisierung der wöchentlichen Berichte mit Python und SQL", 10, False),
    ("Junior-Analyst, Muster AG    06/2018 - 12/2019", 10, True),
    (f"{BULLET} Bereinigung und Prüfung der Verkaufsdaten", 10, False),
    ("AUSBILDUNG", 12, True),
    ("Bachelor in Statistik an der Universität Beispiel    2014 - 2018", 10, False),
    ("KENNTNISSE", 12, True),
    ("Python, SQL, Tableau, Excel", 10, False),
    ("SPRACHEN", 12, True),
    ("Deutsch und Englisch", 10, False),
]

MIXED_FR_EN = [
    (NAME, 20, True),
    ("alex.example@example.com | 555-0100 | Montréal, QC", 10, False),
    ("PROFIL", 12, True),
    ("Analyste de données bilingue avec quatre ans d'expérience.", 10, False),
    ("EXPÉRIENCE PROFESSIONNELLE", 12, True),
    ("Analyste de données chez Exemple Inc.    2020 - aujourd'hui", 10, True),
    (f"{BULLET} Built reporting pipelines in Python and SQL for the finance team", 10, False),
    (f"{BULLET} Automated the weekly dashboards for the sales managers", 10, False),
    (f"{BULLET} Presented the monthly results to the leadership team", 10, False),
    ("FORMATION", 12, True),
    ("Baccalauréat en statistique de l'Université Exemple    2014 - 2018", 10, False),
    (f"{BULLET} Completed a thesis on the forecasting of retail demand", 10, False),
    ("COMPÉTENCES", 12, True),
    ("Python, SQL, Tableau", 10, False),
]


@fixture("pdf_french_resume.pdf")
def _pdf_french() -> bytes:
    return _simple_pdf(FRENCH)


@fixture("pdf_spanish_resume.pdf")
def _pdf_spanish() -> bytes:
    return _simple_pdf(SPANISH)


@fixture("pdf_german_resume.pdf")
def _pdf_german() -> bytes:
    return _simple_pdf(GERMAN)


@fixture("pdf_mixed_fr_en.pdf")
def _pdf_mixed() -> bytes:
    return _simple_pdf(MIXED_FR_EN)


# Logical-order strings (what extraction must return). Arabic is shaped by MuPDF's HTML engine
# and written in logical order; Hebrew is written in VISUAL order (glyphs left to right, as
# older producers do), so its extraction must undo the reordering.
ARABIC_LINES = [
    "مهندس برمجيات",
    "الخبرة المهنية",
    "مهندس برمجيات في شركة مثال من 2019 إلى 2023",
    "تطوير تطبيقات الويب باستخدام بايثون وقواعد البيانات",
    "التعليم",
    "بكالوريوس علوم الحاسوب من جامعة مثال",
    "المهارات",
    "بايثون وقواعد البيانات وتحليل البيانات",
]

HEBREW_LINES = [
    "מהנדס תוכנה",
    "ניסיון תעסוקתי",
    "מהנדס תוכנה בחברת דוגמה",
    "פיתוח יישומי אינטרנט ומסדי נתונים",
    "השכלה",
    "תואר ראשון במדעי המחשב באוניברסיטת דוגמה",
    "כישורים",
    "פיתוח תוכנה וניתוח נתונים",
]


@fixture("pdf_arabic_resume.pdf")
def _pdf_arabic() -> bytes:
    fitz = _fitz()
    doc = fitz.open()
    page = doc.new_page(width=PAGE_W, height=PAGE_H)
    page.insert_text((72, 70), "alex.example@example.com", fontname="helv", fontsize=10)
    body = "".join(f"<p>{line}</p>" for line in ARABIC_LINES)
    page.insert_htmlbox(fitz.Rect(72, 90, 540, 700), f'<div dir="rtl">{body}</div>',
                        css="p {font-size: 13px; margin: 0 0 6px 0;}")
    return _mupdf_bytes(doc)


@fixture("pdf_hebrew_resume.pdf")
def _pdf_hebrew() -> bytes:
    fitz = _fitz()
    font = fitz.Font(script=5)  # MuPDF builtin Noto Serif Hebrew
    doc = fitz.open()
    page = doc.new_page(width=PAGE_W, height=PAGE_H)
    page.insert_font(fontname="heb", fontbuffer=font.buffer)
    page.insert_text((72, 70), "alex.example@example.com", fontname="helv", fontsize=10)
    top = 90.0
    for line in HEBREW_LINES:
        top += 20
        visual = line[::-1]  # visual order: the rightmost (first) letter is drawn last
        width = font.text_length(visual, fontsize=13)
        page.insert_text((540 - width, top), visual, fontname="heb", fontsize=13)
    return _mupdf_bytes(doc)


# ════════════════════════════════════════════════════════════════════════════════════════════
# DOCX
# ════════════════════════════════════════════════════════════════════════════════════════════


@fixture("docx_clean.docx")
def _docx_clean() -> bytes:
    doc = _new_docx()
    _docx_resume_body(doc)
    return _docx_bytes(doc)


@fixture("docx_tables.docx")
def _docx_tables() -> bytes:
    doc = _new_docx()
    doc.add_paragraph(NAME, style="Title")
    doc.add_paragraph(CONTACT)
    doc.add_heading("EXPERIENCE", level=1)
    roles = doc.add_table(rows=2, cols=3)
    for r, cells in enumerate([("Data Analyst", "Example Corp", "Jan 2020 - Present"),
                               ("Junior Analyst", "Sample Industries", "Jun 2018 - Dec 2019")]):
        for col, text in enumerate(cells):
            roles.cell(r, col).text = text
    doc.add_heading("DETAILS", level=1)
    layout = doc.add_table(rows=1, cols=2)  # a layout table: multi-paragraph cells
    left, right = layout.cell(0, 0), layout.cell(0, 1)
    left.text = "SKILLS"
    left.add_paragraph("Python, SQL, Tableau")
    left.add_paragraph("Excel and Statistics")
    right.text = "EDUCATION"
    right.add_paragraph("BSc Statistics, Example University")
    right.add_paragraph("2014 - 2018")
    return _docx_bytes(doc)


def _docx_with_text_box(anchored: bool) -> bytes:
    doc = _new_docx()
    _docx_resume_body(doc)
    anchor = doc.paragraphs[1]  # the contact line anchors the box, as designer templates do
    anchor._p.append(_w(_text_box_run(["CERTIFICATIONS", "Example Certified Data Analyst",
                                        "LANGUAGES", "English and French"], anchored=anchored)))
    return _docx_bytes(doc)


@fixture("docx_text_boxes.docx")
def _docx_text_boxes() -> bytes:
    return _docx_with_text_box(anchored=True)


@fixture("docx_text_box_inline.docx")
def _docx_text_box_inline() -> bytes:
    return _docx_with_text_box(anchored=False)


@fixture("docx_header_footer.docx")
def _docx_header_footer() -> bytes:
    doc = _new_docx()
    doc.add_heading("EXPERIENCE", level=1)
    doc.add_paragraph("Data Analyst, Example Corp    Jan 2020 - Present")
    doc.add_paragraph("Built reporting pipelines in Python and SQL for the finance team", style="List Bullet")
    doc.add_heading("EDUCATION", level=1)
    doc.add_paragraph("BSc Statistics, Example University    2014 - 2018")
    doc.add_page_break()
    doc.add_heading("SKILLS", level=1)
    doc.add_paragraph("Python, SQL, Tableau, Excel, Statistics")
    section = doc.sections[0]
    section.header.paragraphs[0].text = f"{NAME} | alex.example@example.com | 555-0100"
    footer = section.footer.paragraphs[0]
    footer.text = "Portfolio: example.com/alex"
    number = section.footer.add_paragraph("Page ")
    number._p.append(_w(f'<w:fldSimple {_DOCX_NS} w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'))
    return _docx_bytes(doc)


@fixture("docx_hidden_text.docx")
def _docx_hidden_text() -> bytes:
    from docx.enum.style import WD_STYLE_TYPE

    doc = _new_docx()
    _docx_resume_body(doc)
    hidden_style = doc.styles.add_style("Keywords", WD_STYLE_TYPE.CHARACTER)
    hidden_style.font.hidden = True
    para = doc.add_paragraph("References available on request. ")
    para.add_run("Kubernetes Terraform").font.hidden = True  # direct w:vanish
    para.add_run(" Golang Scala", style="Keywords")  # w:vanish inherited from a character style
    return _docx_bytes(doc)


@fixture("docx_white_text.docx")
def _docx_white_text() -> bytes:
    from docx.shared import RGBColor

    doc = _new_docx()
    _docx_resume_body(doc)
    para = doc.add_paragraph("References available on request. ")
    para.add_run("Kubernetes Terraform").font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    para.add_run(" Golang Scala").font.color.rgb = RGBColor(0xFA, 0xFA, 0xFA)
    para.add_run(" Updated October 2026").font.color.rgb = RGBColor(0x80, 0x80, 0x80)  # grey: stays
    return _docx_bytes(doc)


@fixture("docx_tiny_text.docx")
def _docx_tiny_text() -> bytes:
    from docx.shared import Pt

    doc = _new_docx()
    _docx_resume_body(doc)
    para = doc.add_paragraph("References available on request. ")
    para.add_run("Kubernetes Terraform").font.size = Pt(1)
    para.add_run(" Golang Scala").font.size = Pt(1.5)
    para.add_run(" Portfolio at example.com/alex").font.size = Pt(6)  # small but readable: stays
    return _docx_bytes(doc)


@fixture("docx_white_on_dark.docx")
def _docx_white_on_dark() -> bytes:
    from docx.shared import RGBColor

    doc = _new_docx()
    banner = doc.add_table(rows=1, cols=1)
    cell = banner.cell(0, 0)
    cell._tc.get_or_add_tcPr().append(_w(f'<w:shd {_DOCX_NS} w:val="clear" w:color="auto" w:fill="1F3864"/>'))
    run = cell.paragraphs[0].add_run(f"{NAME} - Data Analyst")
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    shaded = doc.add_paragraph()
    shaded._p.get_or_add_pPr().append(_w(f'<w:shd {_DOCX_NS} w:val="clear" w:color="auto" w:fill="000000"/>'))
    shaded.add_run("alex.example@example.com | 555-0100").font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    _docx_resume_body(doc)
    return _docx_bytes(doc)


@fixture("docx_corrupt_zip.docx")
def _docx_corrupt_zip() -> bytes:
    doc = _new_docx()
    _docx_resume_body(doc)
    data = _docx_bytes(doc)
    return data[: len(data) // 2]


@fixture("docx_without_body.docx")
def _docx_without_body() -> bytes:
    """A well-formed word/document.xml whose w:document has no w:body."""
    import re

    doc = _new_docx()
    _docx_resume_body(doc)
    source = zipfile.ZipFile(io.BytesIO(_docx_bytes(doc)))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == "word/document.xml":
                content = re.sub(rb"<w:body>.*</w:body>", b"", content, flags=re.S)
            member = zipfile.ZipInfo(info.filename, date_time=(2020, 1, 1, 0, 0, 0))
            member.compress_type = zipfile.ZIP_DEFLATED  # a ZipInfo ignores the archive's default
            target.writestr(member, content)
    return out.getvalue()


@fixture("docx_zip_without_document.docx")
def _docx_zip_without_document() -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        member = zipfile.ZipInfo("notes/readme.txt", date_time=(2020, 1, 1, 0, 0, 0))
        zf.writestr(member, "Not a Word document.")
    return out.getvalue()


# ════════════════════════════════════════════════════════════════════════════════════════════
# Plain text
# ════════════════════════════════════════════════════════════════════════════════════════════

TEXT_RESUME = "\n".join(text.replace(BULLET, "-") for text, _, _ in RESUME_EN) + "\n"

FRENCH_TEXT = "\n".join(text.replace(BULLET, "-") for text, _, _ in FRENCH) + "\n"


@fixture("text_utf8.txt")
def _text_utf8() -> bytes:
    return TEXT_RESUME.replace("Toronto", "Montréal").encode("utf-8")


@fixture("text_utf8_bom.txt")
def _text_utf8_bom() -> bytes:
    return b"\xef\xbb\xbf" + TEXT_RESUME.encode("utf-8")


@fixture("text_utf16le_bom.txt")
def _text_utf16le() -> bytes:
    return b"\xff\xfe" + TEXT_RESUME.replace("\n", "\r\n").encode("utf-16-le")


@fixture("text_utf16be_bom.txt")
def _text_utf16be() -> bytes:
    return b"\xfe\xff" + TEXT_RESUME.encode("utf-16-be")


@fixture("text_cp1252_french.txt")
def _text_cp1252() -> bytes:
    # cp1252-only bytes (e-acute 0xE9, right single quote 0x92, en dash 0x96) are invalid UTF-8.
    return FRENCH_TEXT.replace("d'expérience", f"d{RSQ}expérience").replace(" - ", f" {chr(0x2013)} ").encode("cp1252")


@fixture("text_two_line_resume.txt")
def _text_two_line() -> bytes:
    return ("\n".join(TWO_LINE_RESUME) + "\n").encode("utf-8")


@fixture("text_binary_named.txt")
def _text_binary_named() -> bytes:
    rng = random.Random(1252)
    return bytes(rng.getrandbits(8) for _ in range(2048))


# ════════════════════════════════════════════════════════════════════════════════════════════
# Type detection (magic bytes first, extension second)
# ════════════════════════════════════════════════════════════════════════════════════════════


@fixture("random_bytes.bin")
def _random_bytes() -> bytes:
    rng = random.Random(42)
    return bytes(rng.getrandbits(8) for _ in range(4096))


@fixture("zero_bytes.pdf")
def _zero_bytes() -> bytes:
    return b""


@fixture("pdf_named_txt.txt")
def _pdf_named_txt() -> bytes:
    return _simple_pdf(RESUME_EN)


@fixture("text_named_pdf.pdf")
def _text_named_pdf() -> bytes:
    return TEXT_RESUME.encode("utf-8")


@fixture("docx_named_pdf.pdf")
def _docx_named_pdf() -> bytes:
    return _docx_clean()


@fixture("gif_image.gif")
def _gif_image() -> bytes:
    from PIL import Image

    return _encode(Image.new("RGB", (64, 32), (200, 30, 30)), "GIF")


@fixture("png_corrupt.png")
def _png_corrupt() -> bytes:
    rng = random.Random(89)
    return b"\x89PNG\r\n\x1a\n" + bytes(rng.getrandbits(8) for _ in range(1024))


# ─── entry point ───────────────────────────────────────────────────────────────────────────────


def main(argv: Sequence[str] = ()) -> int:
    outputs = build_all()
    total = sum(len(data) for data in outputs.values())
    for name, data in outputs.items():
        (FIXTURE_DIR / name).write_bytes(data)
        print(f"{len(data):>9,d}  {name}")
    print(f"{total:>9,d}  total ({len(outputs)} fixtures)")
    if total >= MAX_TOTAL_BYTES:
        print(f"fixtures exceed the {MAX_TOTAL_BYTES:,d}-byte budget", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
