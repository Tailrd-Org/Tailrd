"""Per-format renderers: lay the neutral `Document` out as txt, docx or PDF (EVAL.md section 2.1).

    txt                  plain UTF-8 text
    docx                 python-docx: Title / Heading 1 / List Bullet styles, a table for a side column
    pdf  "latex"         reportlab, Times, `article`-like wide margins, centred header, page numbers
    pdf  "word"          reportlab, Helvetica (a Calibri stand-in), Word's one-inch margins
    pdf  "canva"         reportlab, embedded TrueType face, coloured sidebar with contact and skills
    scan                 any PDF rasterised at 300 dpi, slightly rotated and noisy, image only

Everything is deterministic: reportlab runs with `invariant=1` and fixed metadata, docx archives
are re-packed with fixed member dates, and the scan noise comes from a seeded generator.
reportlab, python-docx, PyMuPDF, Pillow and numpy are imported inside the functions that use them.
"""

from __future__ import annotations

import io
import os
import zipfile
from dataclasses import dataclass
from xml.sax.saxutils import escape

from eval.generate.doc import Document, Item, Section

PAGE = (612.0, 792.0)  # US Letter, points
FIXED_ZIP_DATE = (2000, 1, 1, 0, 0, 0)

BAR_FULL, BAR_EMPTY = "\u25cf", "\u25cb"  # skill bars as text
# Icons live in ZapfDingbats in a PDF (a standard font, so nothing to embed); the text layer of
# such a PDF reads like a real icon-font export, i.e. as stray letters.
ZAPF_ICONS = frozenset(
    "\u2709\u2706\u25c6\u2605\u25cf\u260e"
)  # envelope, phone, diamond, star, dot, desk phone


# --------------------------------------------------------------------------------------------
# txt
# --------------------------------------------------------------------------------------------


def to_txt(doc: Document) -> bytes:
    width = 80
    lines: list[str] = [doc.name]
    if doc.headline:
        lines.append(doc.headline)
    if any(c.icon for c in doc.contacts):
        lines += [f"{c.icon} {c.text}".strip() for c in doc.contacts]
    elif doc.contacts:
        lines.append(" | ".join(c.text for c in doc.contacts))
    for section in doc.sections:
        title = f"{section.icon} {section.title}".strip()
        lines += ["", title, "-" * len(title)]
        for item in section.items:
            lines += _txt_item(item, width)
    return ("\n".join(lines) + "\n").encode("utf-8")


def _txt_item(item: Item, width: int) -> list[str]:
    if item.kind == "entry":
        title, org, dates, place = (item.arg(i) for i in range(4))
        left = title + (f", {org}" if org else "")
        first = left + " " * max(2, width - len(left) - len(dates)) + dates if dates else left
        return [first] + ([place] if place else [])
    if item.kind == "bullet":
        return [f"- {item.arg(0)}"]
    if item.kind in ("para", "skills"):
        return [item.arg(0)]
    if item.kind in ("kv", "group"):
        return [f"{item.arg(0)}: {item.arg(1)}"]
    if item.kind == "bar":
        level = max(0, min(5, int(item.arg(1))))
        return [f"{item.arg(0)}  {BAR_FULL * level}{BAR_EMPTY * (5 - level)}"]
    raise ValueError(f"unknown item {item.kind}")


# --------------------------------------------------------------------------------------------
# docx
# --------------------------------------------------------------------------------------------


def _fixed_zip(data: bytes) -> bytes:
    """The same archive with every member dated 2000-01-01 (python-docx stamps the wall clock)."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            member = zipfile.ZipInfo(info.filename, date_time=FIXED_ZIP_DATE)
            member.compress_type = zipfile.ZIP_DEFLATED
            member.external_attr = 0o600 << 16
            member.create_system = 3
            dst.writestr(member, src.read(info.filename))
    return out.getvalue()


def to_docx(doc: Document) -> bytes:
    import datetime as dt

    import docx
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Inches, Pt

    d = docx.Document()
    fixed = dt.datetime(2000, 1, 1)
    props = d.core_properties
    props.author = props.last_modified_by = "synthetic corpus generator"
    props.title, props.created, props.modified, props.revision = f"Resume of {doc.name}", fixed, fixed, 1
    margin = 0.7 if doc.has_side else 1.0
    section = d.sections[0]
    section.left_margin = section.right_margin = Inches(margin)
    section.top_margin = section.bottom_margin = Inches(0.8 if doc.has_side else 1.0)
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    if doc.running_head:
        section.header.paragraphs[0].text = doc.running_head

    def head(container) -> None:
        container.add_paragraph(doc.name, style="Title")
        if doc.headline:
            container.add_paragraph(doc.headline)
        if any(c.icon for c in doc.contacts):
            for c in doc.contacts:
                container.add_paragraph(f"{c.icon} {c.text}".strip())
        elif doc.contacts:
            container.add_paragraph(" | ".join(c.text for c in doc.contacts))

    def fill(container, sections: list[Section], width_in: float, wide: bool) -> None:
        for sec in sections:
            container.add_paragraph(f"{sec.icon} {sec.title}".strip(), style="Heading 1")
            for item in sec.items:
                if item.kind == "entry":
                    title, org, dates, place = (item.arg(i) for i in range(4))
                    p = container.add_paragraph()
                    p.add_run(title).bold = True
                    if org:
                        p.add_run(f", {org}")
                    if dates and wide:
                        p.paragraph_format.tab_stops.add_tab_stop(Inches(width_in), WD_TAB_ALIGNMENT.RIGHT)
                        p.add_run(f"\t{dates}")
                    elif dates:
                        container.add_paragraph(dates).runs[0].italic = True
                    if place:
                        container.add_paragraph(place).runs[0].italic = True
                elif item.kind == "bullet":
                    container.add_paragraph(item.arg(0), style="List Bullet")
                elif item.kind in ("para", "skills"):
                    container.add_paragraph(item.arg(0))
                elif item.kind in ("kv", "group"):
                    p = container.add_paragraph()
                    p.add_run(f"{item.arg(0)}: ").bold = True
                    p.add_run(item.arg(1))
                elif item.kind == "bar":
                    level = max(0, min(5, int(item.arg(1))))
                    container.add_paragraph(f"{item.arg(0)}  {BAR_FULL * level}{BAR_EMPTY * (5 - level)}")

    if doc.has_side:
        head(d)
        table = d.add_table(rows=1, cols=2)
        table.autofit = False
        left, right = table.rows[0].cells
        left.width, right.width = Inches(2.0), Inches(4.9)
        fill(left, [s for s in doc.sections if s.side], 1.8, wide=False)
        fill(right, [s for s in doc.sections if not s.side], 4.7, wide=True)
        for cell in (left, right):  # a cell starts with an empty paragraph: drop it
            first = cell.paragraphs[0]
            if not first.text and len(cell.paragraphs) > 1:
                first._element.getparent().remove(first._element)
    else:
        head(d)
        fill(d, list(doc.sections), 6.5, wide=True)
    buf = io.BytesIO()
    d.save(buf)
    return _fixed_zip(buf.getvalue())


# --------------------------------------------------------------------------------------------
# pdf
# --------------------------------------------------------------------------------------------


@dataclass(frozen=True)
class PdfLook:
    key: str
    regular: str
    bold: str
    italic: str
    size: float
    leading: float
    name_size: float
    heading_size: float
    margins: tuple[float, float, float, float]  # left, right, top, bottom
    side_margins: tuple[float, float, float, float]  # the same, when there is a side column
    heading_caps: bool
    heading_rule: bool
    heading_color: str
    name_color: str
    centered_header: bool
    bullet_font: str
    page_numbers: bool
    sidebar_width: float
    sidebar_fill: tuple[str, ...] | None  # palette (None: an uncoloured side column)
    always_sidebar: bool
    contact_size: float = 10.0


LATEX = PdfLook(
    key="latex",
    regular="Times-Roman",
    bold="Times-Bold",
    italic="Times-Italic",
    size=11,
    leading=13.8,
    name_size=22,
    heading_size=13,
    margins=(126, 126, 100, 100),
    side_margins=(54, 54, 72, 72),
    heading_caps=False,
    heading_rule=True,
    heading_color="#000000",
    name_color="#000000",
    centered_header=True,
    bullet_font="Symbol",
    page_numbers=True,
    sidebar_width=150,
    sidebar_fill=None,
    always_sidebar=False,
    contact_size=10.5,
)
WORD = PdfLook(
    key="word",
    regular="Helvetica",
    bold="Helvetica-Bold",
    italic="Helvetica-Oblique",
    size=10,
    leading=13,
    name_size=24,
    heading_size=13,
    margins=(72, 72, 72, 72),
    side_margins=(54, 54, 54, 54),
    heading_caps=False,
    heading_rule=False,
    heading_color="#2F5496",
    name_color="#1F3864",
    centered_header=False,
    bullet_font="Symbol",
    page_numbers=False,
    sidebar_width=160,
    sidebar_fill=None,
    always_sidebar=False,
    contact_size=10,
)
CANVA = PdfLook(
    key="canva",
    regular="Vera",
    bold="Vera-Bold",
    italic="Vera-Italic",
    size=9,
    leading=12.5,
    name_size=26,
    heading_size=11,
    margins=(0, 34, 40, 40),
    side_margins=(0, 34, 40, 40),
    heading_caps=True,
    heading_rule=False,
    heading_color="#223A5E",
    name_color="#223A5E",
    centered_header=False,
    bullet_font="Vera",
    page_numbers=False,
    sidebar_width=206,
    sidebar_fill=("#223A5E", "#1F4E5F", "#4A2C5A", "#24523B", "#3A4750", "#6B2737"),
    always_sidebar=True,
    contact_size=8,
)
LOOKS = {"latex": LATEX, "word": WORD, "canva": CANVA}

_FONTS_READY = False


def _ensure_fonts() -> None:
    """Bitstream Vera ships with reportlab; it is the embedded face of the canva-like look."""
    global _FONTS_READY
    if _FONTS_READY:
        return
    import reportlab
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    folder = os.path.join(os.path.dirname(reportlab.__file__), "fonts")
    for name, file in (
        ("Vera", "Vera.ttf"),
        ("Vera-Bold", "VeraBd.ttf"),
        ("Vera-Italic", "VeraIt.ttf"),
        ("Vera-BoldItalic", "VeraBI.ttf"),
    ):
        pdfmetrics.registerFont(TTFont(name, os.path.join(folder, file)))
    pdfmetrics.registerFontFamily(
        "Vera", normal="Vera", bold="Vera-Bold", italic="Vera-Italic", boldItalic="Vera-BoldItalic"
    )
    _FONTS_READY = True


def _text(value: str) -> str:
    """Markup-safe text that every font of the looks can show (Latin-1 plus the usual punctuation)."""
    return escape(value.encode("cp1252", "replace").decode("cp1252"))


def _icon(char: str) -> str:
    return f'<font name="ZapfDingbats">{char}</font>&nbsp;' if char in ZAPF_ICONS else ""


class _Styles:
    """The paragraph styles of one look, in the colours of the main column or of a tinted sidebar."""

    def __init__(self, look: PdfLook, text: str, heading: str, rule: str) -> None:
        from reportlab.lib.colors import HexColor
        from reportlab.lib.enums import TA_CENTER, TA_RIGHT
        from reportlab.lib.styles import ParagraphStyle

        self.look, self.rule = look, HexColor(rule)
        base = dict(fontName=look.regular, fontSize=look.size, leading=look.leading, textColor=HexColor(text))
        self.body = ParagraphStyle("body", **base, spaceAfter=2.5)
        self.right = ParagraphStyle("right", **base, alignment=TA_RIGHT)
        self.bullet = ParagraphStyle(
            "bullet",
            **base,
            leftIndent=13,
            bulletIndent=2,
            bulletFontName=look.bullet_font,
            bulletFontSize=look.size,
            spaceAfter=1.5,
        )
        self.line = ParagraphStyle("line", **base, spaceAfter=1)
        self.heading = ParagraphStyle(
            "heading",
            fontName=look.bold,
            fontSize=look.heading_size,
            leading=look.heading_size + 3,
            textColor=HexColor(heading),
            spaceBefore=11,
            spaceAfter=3,
            keepWithNext=1,
        )
        align = TA_CENTER if look.centered_header else 0
        self.name = ParagraphStyle(
            "name",
            fontName=look.bold,
            fontSize=look.name_size,
            leading=look.name_size + 4,
            textColor=HexColor(look.name_color),
            alignment=align,
            spaceAfter=2,
        )
        self.headline = ParagraphStyle(
            "headline",
            **{**base, "fontSize": look.size + 1.5, "leading": look.leading + 2},
            alignment=align,
            spaceAfter=4,
        )
        contact_base = {**base, "fontSize": look.contact_size, "leading": look.contact_size + 3.5}
        self.contact = ParagraphStyle(
            "contact", **contact_base, alignment=align, spaceAfter=1.5, splitLongWords=0
        )


def _skill_bar_flowable():
    from reportlab.lib.colors import HexColor
    from reportlab.platypus import Flowable

    class SkillBar(Flowable):
        def __init__(self, label: str, level: int, width: float, st: _Styles, fill: str, track: str) -> None:
            super().__init__()
            self.label, self.level, self.avail, self.st = label, level, width, st
            self.fill, self.track = HexColor(fill), HexColor(track)

        def wrap(self, aw, ah):
            self.width = min(self.avail, aw)
            return self.width, self.st.look.size + 15

        def draw(self):
            c, look = self.canv, self.st.look
            c.setFont(look.regular, look.size)
            c.setFillColor(self.st.body.textColor)
            c.drawString(0, 13, self.label)  # the label, then its bar underneath, then some air
            seg = (self.width - 4 * 3) / 5
            for k in range(5):
                c.setFillColor(self.fill if k < self.level else self.track)
                c.rect(k * (seg + 3), 5, seg, 4, stroke=0, fill=1)

    return SkillBar


def _entry(item: Item, st: _Styles, width: float, wide: bool) -> list:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle

    title, org, dates, place = (item.arg(i) for i in range(4))
    left = f"<b>{_text(title)}</b>" + (f", {_text(org)}" if org else "")
    if wide and dates:
        dates_w = pdfmetrics.stringWidth(_text(dates), st.look.regular, st.look.size) + 10
        rows = [[Paragraph(left, st.body), Paragraph(_text(dates), st.right)]]
        if place:
            rows.append([Paragraph(f"<i>{_text(place)}</i>", st.body), ""])
        table = Table(rows, colWidths=[width - dates_w, dates_w], hAlign="LEFT")
        table.setStyle(
            TableStyle(
                [("VALIGN", (0, 0), (-1, -1), "TOP")]
                + [(f"{side}PADDING", (0, 0), (-1, -1), 0) for side in ("LEFT", "RIGHT", "TOP", "BOTTOM")]
            )
        )
        return [Spacer(1, 3), table]
    out = [Spacer(1, 3), Paragraph(left, st.body)]
    for extra in (dates, place):
        if extra:
            out.append(Paragraph(f"<i>{_text(extra)}</i>", st.body))
    return out


def _section_flowables(
    sec: Section, st: _Styles, width: float, wide: bool, bar_colors: tuple[str, str], *, keep: bool = True
) -> list:
    from reportlab.platypus import HRFlowable, KeepTogether, Paragraph

    look = st.look
    title = sec.title.upper() if look.heading_caps else sec.title
    out: list = [Paragraph(_icon(sec.icon) + _text(title), st.heading)]
    if look.heading_rule:
        rule = HRFlowable(width="100%", thickness=0.6, color=st.rule, spaceBefore=0, spaceAfter=4)
        rule.keepWithNext = 1
        out.append(rule)
    items = list(sec.items)
    i = 0
    while i < len(items):
        item = items[i]
        if item.kind == "entry":
            block = _entry(item, st, width, wide)
            if i + 1 < len(items) and items[i + 1].kind == "bullet":  # never strand an entry at a page end
                block.append(Paragraph(_text(items[i + 1].arg(0)), st.bullet, bulletText="\u2022"))
                i += 1
            out += [KeepTogether(block)] if keep else block
        elif item.kind == "bullet":
            out.append(Paragraph(_text(item.arg(0)), st.bullet, bulletText="\u2022"))
        elif item.kind == "para":
            out.append(Paragraph(_text(item.arg(0)), st.body))
        elif item.kind in ("kv", "group"):
            out.append(Paragraph(f"<b>{_text(item.arg(0))}:</b> {_text(item.arg(1))}", st.body))
        elif item.kind == "skills":
            if wide:
                out.append(Paragraph(_text(item.arg(0)), st.body))
            else:
                out += [Paragraph(_text(s.strip()), st.line) for s in item.arg(0).split(",")]
        elif item.kind == "bar":
            out.append(_skill_bar_flowable()(item.arg(0), int(item.arg(1)), width, st, *bar_colors))
        i += 1
    return out


def _header_flowables(doc: Document, st: _Styles, with_contacts: bool) -> list:
    from reportlab.platypus import Paragraph

    out = [Paragraph(_text(doc.name), st.name)]
    if doc.headline:
        out.append(Paragraph(_text(doc.headline), st.headline))
    if with_contacts and doc.contacts:
        out += _contact_flowables(doc, st)
    return out


def _contact_flowables(doc: Document, st: _Styles) -> list:
    from reportlab.platypus import Paragraph

    if any(c.icon for c in doc.contacts):
        return [Paragraph(_icon(c.icon) + _text(c.text), st.contact) for c in doc.contacts]
    return [Paragraph("  |  ".join(_text(c.text) for c in doc.contacts), st.contact)]


def to_pdf(doc: Document, look_key: str, *, accent_roll: int = 0, title: str = "") -> bytes:
    """`accent_roll` picks the sidebar colour of the canva-like look from its palette."""
    from reportlab.lib.colors import HexColor
    from reportlab.platypus import (
        BaseDocTemplate,
        Frame,
        FrameBreak,
        KeepInFrame,
        NextPageTemplate,
        PageTemplate,
    )

    _ensure_fonts()
    look = LOOKS[look_key]
    if look.always_sidebar and not doc.has_side:
        doc = doc.with_sidebar(frozenset({"skills", "certifications"}))
    sidebar = doc.has_side
    pw, ph = PAGE
    left, right, top, bottom = look.side_margins if sidebar else look.margins
    fill = look.sidebar_fill[accent_roll % len(look.sidebar_fill)] if look.sidebar_fill else None

    main = _Styles(look, "#222222" if fill else "#000000", fill or look.heading_color, "#444444")
    if fill:
        main.name.textColor = HexColor(fill)
    side_st = _Styles(look, "#FFFFFF", "#FFFFFF", "#FFFFFF") if fill else main

    def no_pad(x, y, w, h, fid) -> Frame:
        return Frame(x, y, w, h, id=fid, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)

    def decorate(canvas, _doc) -> None:
        canvas.saveState()
        if fill:
            canvas.setFillColor(HexColor(fill))
            canvas.rect(0, 0, look.sidebar_width, ph, stroke=0, fill=1)
        canvas.setFont(look.regular, 9)
        canvas.setFillColor(HexColor("#555555"))
        if look.page_numbers:
            canvas.drawCentredString(pw / 2, bottom - 30, str(canvas.getPageNumber()))
        if doc.running_head:
            canvas.drawRightString(
                pw - right, ph - 34, f"{_text(doc.running_head)} - Page {canvas.getPageNumber()}"
            )
        canvas.restoreState()

    buf = io.BytesIO()
    pdf = BaseDocTemplate(
        buf,
        pagesize=PAGE,
        invariant=1,
        pageCompression=1,
        title=title or f"Resume of {doc.name}",
        author="synthetic corpus generator",
        subject="synthetic resume",
        creator="eval.generate",
        initialFontName=look.regular,
        initialFontSize=look.size,
    )
    ph_inner = ph - top - bottom
    if sidebar:
        side_x = 18 if fill else left
        side_w = look.sidebar_width - 36 if fill else look.sidebar_width
        main_x = (look.sidebar_width + 24) if fill else (left + look.sidebar_width + 18)
        main_w = pw - main_x - right
        side_frame = no_pad(side_x, bottom, side_w, ph_inner, "side")
        first = PageTemplate(
            id="first",
            frames=[side_frame, no_pad(main_x, bottom, main_w, ph_inner, "main1")],
            onPage=decorate,
        )
        later = PageTemplate(
            id="later", frames=[no_pad(main_x, bottom, main_w, ph_inner, "main2")], onPage=decorate
        )
        pdf.addPageTemplates([first, later])
        bar_colors = ("#FFFFFF", "#7C8AA3") if fill else (look.heading_color, "#CCCCCC")
        side_flow: list = []
        if fill:
            side_flow += _contact_flowables(doc, side_st)
        for sec in doc.sections:
            if sec.side:
                # KeepInFrame draws its content directly, so nothing in it may need splitting
                side_flow += _section_flowables(sec, side_st, side_w, False, bar_colors, keep=False)
        story: list = [NextPageTemplate("later"), KeepInFrame(side_w, ph_inner, side_flow, mode="shrink")]
        story.append(FrameBreak())
        story += _header_flowables(doc, main, with_contacts=not fill)
        for sec in doc.sections:
            if not sec.side:
                story += _section_flowables(sec, main, main_w, wide=True, bar_colors=bar_colors)
    else:
        width = pw - left - right
        pdf.addPageTemplates(
            [PageTemplate(id="main", frames=[no_pad(left, bottom, width, ph_inner, "main")], onPage=decorate)]
        )
        story = _header_flowables(doc, main, with_contacts=True)
        for sec in doc.sections:
            story += _section_flowables(
                sec, main, width, wide=True, bar_colors=(look.heading_color, "#CCCCCC")
            )
    pdf.build(story)
    return buf.getvalue()


# --------------------------------------------------------------------------------------------
# scan
# --------------------------------------------------------------------------------------------


def scan(pdf_bytes: bytes, seed: int) -> bytes:
    """Rasterise every page at 300 dpi (grey), rotate it by at most half a degree, add sensor
    noise and a few specks, and embed the result as an image-only PDF."""
    import numpy as np
    import pymupdf
    from PIL import Image
    from reportlab import rl_config
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    rng = np.random.default_rng(seed)
    source = pymupdf.open(stream=pdf_bytes, filetype="pdf")
    out = io.BytesIO()
    use_a85 = rl_config.useA85
    rl_config.useA85 = 0  # binary-safe streams: a JPEG stays a JPEG
    try:
        pdf = canvas.Canvas(out, pagesize=PAGE, invariant=1, pageCompression=1)
        pdf.setTitle("Scanned resume")
        for page in source:
            pix = page.get_pixmap(dpi=300, colorspace=pymupdf.csGRAY, alpha=False)
            image = Image.frombytes("L", (pix.width, pix.height), pix.samples)
            angle = float(rng.uniform(-0.5, 0.5))
            image = image.rotate(angle, resample=Image.Resampling.BILINEAR, fillcolor=255)
            pixels = np.asarray(image, dtype=np.float32)
            pixels += rng.standard_normal(pixels.shape, dtype=np.float32) * 4.0
            ys = rng.integers(0, pix.height - 3, 220)
            xs = rng.integers(0, pix.width - 3, 220)
            for y, x, tone in zip(ys, xs, rng.integers(60, 150, 220)):
                pixels[y : y + 2, x : x + 2] = tone
            noisy = Image.fromarray(np.clip(pixels, 0, 255).astype(np.uint8), "L")
            jpeg = io.BytesIO()
            noisy.save(jpeg, "JPEG", quality=58, dpi=(300, 300))
            pdf.setPageSize((page.rect.width, page.rect.height))
            pdf.drawImage(
                ImageReader(io.BytesIO(jpeg.getvalue())), 0, 0, width=page.rect.width, height=page.rect.height
            )
            pdf.showPage()
        pdf.save()
    finally:
        rl_config.useA85 = use_a85
    return out.getvalue()
